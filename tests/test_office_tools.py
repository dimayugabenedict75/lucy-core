"""P3 Office tools: docx_create/edit, xlsx_create/edit, pptx_create/edit, pdf_create, office_read, office_render.
Real files, real libraries (python-docx, openpyxl, python-pptx, reportlab, pypdf, pypdfium2). Tests that need LibreOffice
skip themselves when it is not installed (it is optional)."""
import asyncio, hashlib, os, pathlib, shutil, sys, tempfile
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "src"))
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from lucy import paths as _paths  # noqa: E402
from lucy.tools import registry as R, routing, _office as O, pdf_tools, xlsx_tools  # noqa: E402
from lucy.safety import policy  # noqa: E402
from lucy.safety.taint import unwrap  # noqa: E402
from _safety_helpers import Env  # noqa: E402

reg = R.Registry(); reg.discover("lucy.tools")
HAVE_LO = O.find_soffice() is not None


class proj:
    """SAFE_ROOT = ROOT = a temp folder, RUNTIME = <tmp>/runtime."""
    def __enter__(self):
        self.saved = (_paths.SAFE_ROOT, _paths.ROOT, _paths.RUNTIME)
        self.tmp = pathlib.Path(tempfile.mkdtemp()).resolve()
        _paths.SAFE_ROOT = _paths.ROOT = self.tmp
        _paths.RUNTIME = self.tmp / "runtime"
        return self.tmp

    def __exit__(self, *a):
        from lucy import safety
        safety.set_default(None)                       # the default layer was built on THIS temp folder: never leave it for other test files
        _paths.SAFE_ROOT, _paths.ROOT, _paths.RUNTIME = self.saved
        shutil.rmtree(self.tmp, ignore_errors=True)


def call(name, **args):
    return asyncio.run(reg.run(name, args))


def png(path, size=(320, 200), color=(40, 90, 160), fmt="PNG"):
    from PIL import Image
    pathlib.Path(path).parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, color).save(path, fmt)
    return pathlib.Path(path)


def sha(p):
    return hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()


# ------------------------------------------------------------------------------------------------ registration
def test_all_office_tools_are_registered_with_small_model_shapes():
    want = {"docx_create": "write", "docx_edit": "write", "xlsx_create": "write", "xlsx_edit": "write", "pptx_create": "write",
            "pptx_edit": "write", "pdf_create": "write", "office_read": "read", "office_render": "read"}
    for name, tier in want.items():
        s = reg.get(name)
        assert s is not None and s.tier == tier, name
        fn = s.schema["function"]
        assert len(fn["description"].split(". ")) <= 2 and len(fn["description"]) < 330, name
        assert len(fn["parameters"]["properties"]) <= 8, name
        assert "path" in fn["parameters"]["properties"] and "path" in fn["parameters"]["required"], name
    assert "path" in policy.PATH_KEYS          # the safety layer sees the target file (rules, snapshot, approval) through `path`


# ------------------------------------------------------------------------------------------------ delivery
def test_a_finished_file_is_delivered_with_a_media_line_relative_to_the_safe_root():
    with proj() as t:
        assert call("docx_create", path=str(t / "r.docx"), markdown="# Hi\n\nText.") == "MEDIA:r.docx"
        assert call("docx_create", path=str(t / "sub" / "r2"), markdown="# Hi") == "MEDIA:sub/r2.docx"     # folder made, extension added


def test_a_path_with_spaces_is_delivered_as_a_copy_that_a_media_line_can_carry():
    with proj() as t:
        out = call("docx_create", path=str(t / "My Docs" / "Q3 report.docx"), markdown="# Hi")
        assert out.startswith("MEDIA:") and " " not in out
        assert (t / "My Docs" / "Q3 report.docx").is_file()                              # the real file is where the user asked
        delivered = pathlib.Path(_paths.SAFE_ROOT) / out[len("MEDIA:"):]
        assert delivered.is_file() and delivered.name == "Q3_report.docx"


def test_problems_come_back_as_text_the_model_can_read_and_the_file_is_not_delivered_yet():
    with proj() as t:
        out = call("pptx_create", path=str(t / "d.pptx"), slides=[{"title": "Too much", "bullets": ["word " * 40] * 14}])
        assert not out.startswith("MEDIA:") and "NOT sent to the chat yet" in out and "Slide 1" in out and "too much text" in out
        assert (t / "d.pptx").is_file()                                                  # saved, only the delivery waits for the fix


# ------------------------------------------------------------------------------------------------ Word
MD = ("# Q3 Report\n\nIntro with **bold** and *italic*.\n\n## Numbers\n\n| Item | Qty |\n|---|---|\n| Pens | 4 |\n| Books | 10 |\n\n"
      "- first\n- second\n  - nested\n\n1. one\n2. two\n\n[pagebreak]\n\n## Next\n\n1. again\n2. more\n")


def test_docx_markdown_round_trip():
    with proj() as t:
        assert call("docx_create", path=str(t / "r.docx"), markdown=MD).startswith("MEDIA:")
        txt = call("office_read", path=str(t / "r.docx"))
        for want in ["# Q3 Report", "## Numbers", "Intro with bold and italic.", "| Item | Qty |", "| Pens | 4 |", "- first", "  - nested",
                     "1. one", "[page break]", "## Next"]:
            assert want in txt, want
        import docx
        d = docx.Document(str(t / "r.docx"))
        assert d.core_properties.title == "Q3 Report" and d.core_properties.author == "Lucy"
        ids = {p.text: p._p.pPr.numPr.numId.val for p in d.paragraphs if p.style.name == "List Number"}
        assert ids["one"] == ids["two"] and ids["again"] == ids["more"] and ids["one"] != ids["again"]   # each list is its own, starting at 1


def test_docx_blocks_accept_aliases_json_text_and_rows_of_objects():
    with proj() as t:
        blocks = [{"type": "h1", "text": "Hi"}, {"type": "paragraph", "text": "x"}, {"type": "list", "items": ["a", "b"]},
                  {"type": "table", "rows": [{"a": 1, "b": 2}, {"a": 3, "b": 4}]}, {"type": "page_break"}]
        assert call("docx_create", path=str(t / "a.docx"), blocks=blocks).startswith("MEDIA:")
        txt = call("office_read", path=str(t / "a.docx"))
        assert "## Hi" in txt and "- a" in txt and "| a | b |" in txt and "| 1 | 2 |" in txt
        import json
        assert call("docx_create", path=str(t / "b.docx"), blocks=json.dumps(blocks)).startswith("MEDIA:")      # double-encoded list


def test_docx_resume_style_is_compact_and_caps_the_section_headings():
    with proj() as t:
        call("docx_create", path=str(t / "cv.docx"), style="resume", markdown="# Jane Doe\n\nManila | jane@example.com\n\n## Experience\n- Built things\n")
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        d = docx.Document(str(t / "cv.docx"))
        assert d.sections[0].left_margin.inches < 0.8 and d.styles["Heading 1"].font.all_caps is True
        assert d.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.CENTER                    # the contact line under the name
        call("docx_create", path=str(t / "a4.docx"), paper="a4", markdown="# x")
        assert abs(docx.Document(str(t / "a4.docx")).sections[0].page_width.inches - 8.27) < 0.02


def test_docx_errors_say_what_is_wrong_and_how_to_fix_it():
    with proj() as t:
        p = str(t / "x.docx")
        assert call("docx_create", path=p, blocks=[{"type": "chart"}]).startswith("Error: block 1: unknown type 'chart'. Use one of")
        assert call("docx_create", path=p, blocks=[{"type": "para", "text": " "}]).startswith("Error: block 1 (para): 'text' is empty")
        assert call("docx_create", path=p, blocks=[{"type": "table", "rows": []}]).startswith("Error: block 1: a table needs 'rows'")
        assert call("docx_create", path=p).startswith("Error: nothing to write")
        assert call("docx_create", path=str(t / "x.txt"), markdown="# a").startswith("Error: a Word file must end with .docx")
        assert call("docx_create", path="", markdown="# a").startswith("Error: 'path' is missing")
        assert call("docx_create", path=p, markdown="![x](nope.png)").startswith("Error: block 1: image not found")
        assert not (t / "x.docx").exists()                                               # nothing half-written


def test_docx_images_are_embedded_and_unusual_formats_converted():
    with proj() as t:
        png(t / "a.png"); png(t / "b.webp", fmt="WEBP")
        assert call("docx_create", path=str(t / "i.docx"), markdown=f"![First]({(t / 'a.png').as_posix()})\n\n![Second]({(t / 'b.webp').as_posix()})").startswith("MEDIA:")
        import docx
        assert len(docx.Document(str(t / "i.docx")).inline_shapes) == 2
        assert "[image]" in call("office_read", path=str(t / "i.docx"))
        (t / "fake.png").write_text("not an image")
        assert "not a usable image" in call("docx_create", path=str(t / "j.docx"), markdown=f"![x]({(t / 'fake.png').as_posix()})")


def test_docx_edit_ops_change_the_file_and_read_back():
    with proj() as t:
        p = str(t / "r.docx")
        call("docx_create", path=p, markdown=MD)
        out = call("docx_edit", path=p, ops=[
            {"op": "replace", "find": "Intro with", "replace": "Opening with"},
            {"op": "insert_after", "match": "Numbers", "markdown": "Inserted **after** the heading."},
            {"op": "insert_before", "match": "## Next".replace("## ", ""), "markdown": "Before next."},
            {"op": "set_cell", "table": 1, "row": 2, "col": 2, "text": "5"},
            {"op": "add_row", "table": 1, "cells": ["Pencils", "3"]},
            {"op": "delete", "match": "Before next"},
            {"op": "append", "markdown": "## Appendix\n- x\n- y"}])
        assert out.startswith("MEDIA:"), out
        txt = call("office_read", path=p)
        assert "Opening with bold" in txt and "Intro with" not in txt
        assert txt.index("Numbers") < txt.index("Inserted after the heading.") < txt.index("| Item | Qty |")
        assert "| Pens | 5 |" in txt and "| Pencils | 3 |" in txt and "Before next" not in txt and "## Appendix" in txt and txt.rstrip().endswith("- y")


def test_docx_replace_works_across_formatting_runs_and_in_tables():
    with proj() as t:
        p = str(t / "r.docx")
        call("docx_create", path=p, markdown="Total is **$100** today.\n\n| Name | Note |\n|---|---|\n| Bob | owes $100 |\n")
        # the match spans three runs (plain, bold, plain): it must still work
        assert call("docx_edit", path=p, ops=[{"op": "replace", "find": "is $100 today", "replace": "is $250 now"}]).startswith("MEDIA:")
        assert "Total is $250 now." in call("office_read", path=p)
        assert call("docx_edit", path=p, ops=[{"op": "replace", "find": "$100", "replace": "$300"}]).startswith("MEDIA:")      # table cell
        txt = call("office_read", path=p)
        assert "owes $300" in txt and "$100" not in txt


def test_docx_edit_is_all_or_nothing_and_errors_are_specific():
    with proj() as t:
        p = str(t / "r.docx")
        call("docx_create", path=p, markdown=MD)
        before = sha(p)
        out = call("docx_edit", path=p, ops=[{"op": "replace", "find": "Pens", "replace": "Pencils"}, {"op": "replace", "find": "NOPE-NOPE", "replace": "x"}])
        assert out.startswith("Error: op 2: 'NOPE-NOPE' was not found") and sha(p) == before         # op 1 was rolled back too
        assert call("docx_edit", path=p, ops=[{"op": "set_cell", "table": 3, "row": 1, "col": 1, "text": "x"}]).startswith("Error: op 1: table 3 does not exist")
        assert "inside a table" in call("docx_edit", path=p, ops=[{"op": "insert_after", "match": "Books", "markdown": "x"}])
        assert call("docx_edit", path=p, ops=[{"op": "explode"}]).startswith("Error: op 1: unknown op 'explode'")
        assert call("docx_edit", path=p, ops=[]).startswith("Error: 'ops' is empty")
        assert call("docx_edit", path=str(t / "missing.docx"), ops=[{"op": "append", "markdown": "x"}]).startswith("Error: file not found")
        (t / "bad.docx").write_text("this is not a zip")
        assert "could not be opened as a Word file" in call("docx_edit", path=str(t / "bad.docx"), ops=[{"op": "append", "markdown": "x"}])


# ------------------------------------------------------------------------------------------------ Excel
def test_xlsx_create_styles_formulas_formats_and_cleans_values():
    with proj() as t:
        p = str(t / "b.xlsx")
        out = call("xlsx_create", path=p, sheets=[{
            "name": "Budget", "columns": ["Item", "Qty", "Price", "Total"],
            "rows": [["Pens", "4", "1.50", "=B2*C2"], ["Books", "10", "1,200.00", "=B3*C3"], ["Start", "2025-01-31", "007", 0]],
            "formulas": {"D5": "SUM(D2:D4)"}, "number_formats": {"C": "currency", "D2:D5": "peso"}, "column_widths": {"A": 30}}])
        assert out.startswith("MEDIA:"), out
        import openpyxl, datetime
        ws = openpyxl.load_workbook(p)["Budget"]
        assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("1F3864") and ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:D4"
        assert ws["B2"].value == 4 and ws["C2"].value == 1.5 and ws["C3"].value == 1200 and ws["C4"].value == "007"   # numbers fixed, ids kept
        assert ws["B4"].value == datetime.datetime(2025, 1, 31) and ws["B4"].number_format == "yyyy-mm-dd"
        assert ws["D2"].value == "=B2*C2" and ws["D5"].value == "=SUM(D2:D4)"                                          # formulas, not results
        assert ws["C2"].number_format == '"$"#,##0.00' and "\u20b1" in ws["D3"].number_format and ws.column_dimensions["A"].width == 30
        assert ws["A2"].font.name == "Arial"


def test_xlsx_sheet_names_are_made_valid_and_unique_and_top_level_rows_are_forgiven():
    with proj() as t:
        p = str(t / "s.xlsx")
        assert call("xlsx_create", path=p, sheets=[{"name": "Q3/Q4: [Sales]?", "columns": ["a"], "rows": [[1]]}, {"name": "q3 q4  sales", "columns": ["a"], "rows": [[1]]}]).startswith("MEDIA:")
        import openpyxl
        names = openpyxl.load_workbook(p).sheetnames
        assert len(set(n.lower() for n in names)) == 2 and all(not set("[]:*?/\\") & set(n) for n in names)
        assert call("xlsx_create", path=str(t / "f.xlsx"), columns=["x", "y"], rows=[[1, 2]], sheets=None).startswith("MEDIA:")
        assert call("xlsx_create", path=str(t / "g.xlsx"), sheets=[{"name": "E"}]).startswith("Error: sheet 'E' is empty")
        assert call("xlsx_create", path=str(t / "h.xlsx")).startswith("Error: 'sheets' is missing")


def test_xlsx_edit_ops_and_errors():
    with proj() as t:
        p = str(t / "b.xlsx")
        call("xlsx_create", path=p, sheets=[{"name": "Budget", "columns": ["Item", "Qty", "Price", "Total"], "rows": [["Pens", 4, 1.5, "=B2*C2"]]}, {"name": "Notes", "columns": ["n"], "rows": [["a"]]}])
        out = call("xlsx_edit", path=p, ops=[
            {"op": "set_cells", "sheet": "Budget", "cells": {"B2": "9", "A3": "Books"}},
            {"op": "append_rows", "sheet": 1, "rows": [["Pencils", 3, 0.5, "=B4*C4"]]},
            {"op": "number_format", "sheet": 1, "range": "C", "format": "currency"},
            {"op": "rename_sheet", "sheet": 2, "name": "Memo"},
            {"op": "add_sheet", "name": "Extra", "columns": ["k"], "rows": [[1]]},
            {"op": "delete_sheet", "sheet": "Extra"}])
        assert out.startswith("MEDIA:"), out
        import openpyxl
        wb = openpyxl.load_workbook(p)
        ws = wb["Budget"]
        assert wb.sheetnames == ["Budget", "Memo"] and ws["B2"].value == 9 and ws["A3"].value == "Books"
        assert [c.value for c in ws[4]] == ["Pencils", 3, 0.5, "=B4*C4"] and ws["C2"].number_format == '"$"#,##0.00' and ws["D2"].value == "=B2*C2"
        before = sha(p)
        assert call("xlsx_edit", path=p, ops=[{"op": "set_cells", "sheet": "Nope", "cells": {"A1": 1}}]) == "Error: op 1: no sheet 'Nope'. The sheets are: Budget, Memo."
        assert call("xlsx_edit", path=p, ops=[{"op": "set_cells", "cells": {"A1": "changed"}}, {"op": "set_cells", "cells": {"ZZZZ99999x": 1}}]).startswith("Error: op 2")
        assert sha(p) == before                                                          # atomic
        assert call("xlsx_edit", path=p, ops=[{"op": "delete_sheet", "sheet": 1}, {"op": "delete_sheet", "sheet": 1}]).startswith("Error: op 2: a workbook needs at least one sheet")


def test_xlsx_edit_refuses_to_silently_drop_charts():
    with proj() as t:
        import openpyxl
        from openpyxl.chart import BarChart, Reference
        wb = openpyxl.Workbook(); ws = wb.active
        for r in [["k", "v"], ["a", 1], ["b", 2]]:
            ws.append(r)
        ch = BarChart(); ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True); ws.add_chart(ch, "D2")
        p = str(t / "c.xlsx"); wb.save(p)
        before = sha(p)
        out = call("xlsx_edit", path=p, ops=[{"op": "set_cells", "cells": {"B2": 5}}])
        assert out.startswith("Error:") and "allow_chart_loss" in out and sha(p) == before
        assert call("xlsx_edit", path=p, ops=[{"op": "set_cells", "cells": {"B2": 5}}], allow_chart_loss=True).startswith("MEDIA:")


def test_xlsx_formula_errors_are_reported_before_delivery():
    if not HAVE_LO:
        return                                                                           # LibreOffice is optional
    with proj() as t:
        out = call("xlsx_create", path=str(t / "e.xlsx"), sheets=[{"name": "S", "columns": ["a", "b"], "rows": [[1, 0], ["=A2/B2", "=NOPE(1)"]]}])
        assert not out.startswith("MEDIA:") and "#DIV/0!" in out and "#NAME?" in out and "S!A3" in out


def test_office_read_shows_formulas_with_their_values_when_libreoffice_can_calculate():
    with proj() as t:
        p = str(t / "b.xlsx")
        call("xlsx_create", path=p, sheets=[{"name": "B", "columns": ["Qty", "Price", "Total"], "rows": [[4, 2.5, "=A2*B2"]], "number_formats": {"B": "percent"}}])
        txt = call("office_read", path=p)
        assert "=A2*B2" in txt and "## Sheet \"B\" A1:C2" in txt and "freeze=A2" in txt
        if HAVE_LO:
            assert "=A2*B2 \u2192 10" in txt
        assert "250%" in txt                                                              # percent format shown as a percent
        assert call("office_read", path=p, sheet="nope").startswith("Error: no sheet 'nope'")


# ------------------------------------------------------------------------------------------------ PowerPoint
DECK = """# Q3 Review
Sales, wins and next steps
Notes: Welcome everyone.

---
## Highlights
- Revenue up **12%**
- Three new customers
  - Acme
Notes: Mention Acme.

---
## Regions
| Region | Revenue |
|---|---|
| North | $1,200 |
| South | $900 |

---
## Thank you
Questions?
"""


def test_pptx_markdown_deck_round_trip_with_notes_tables_and_slide_numbers():
    with proj() as t:
        p = str(t / "d.pptx")
        assert call("pptx_create", path=p, markdown=DECK) == "MEDIA:d.pptx"
        txt = call("office_read", path=p)
        assert "Deck: 4 slide(s)" in txt and "--- Slide 1 (Title Slide) ---" in txt and "Title: Q3 Review" in txt and "Sales, wins and next steps" in txt
        assert "Revenue up 12%" in txt and "  - Acme" in txt and "Notes: Mention Acme." in txt and "| North | $1,200 |" in txt and "Title: Thank you" in txt
        import pptx
        prs = pptx.Presentation(p)
        assert abs(prs.slide_width.inches - 13.333) < 0.01 and prs.core_properties.title == "Q3 Review"
        assert any(sh.name == "Slide Number" for sh in prs.slides[1].shapes) and not any(sh.name == "Slide Number" for sh in prs.slides[0].shapes)
        assert call("office_read", path=p, pages="2")[:20] == "Deck: 4 slide(s), sh" and "Thank you" not in call("office_read", path=p, pages="2")


def test_pptx_layouts_are_fixed_to_widescreen_so_new_slides_in_powerpoint_look_right():
    with proj() as t:
        p = str(t / "d.pptx"); call("pptx_create", path=p, markdown=DECK)
        import pptx
        prs = pptx.Presentation(p)
        names = [l.name for l in prs.slide_layouts]
        assert "Title and Content" in names and "Comparison" not in names                # 4:3-only layouts removed
        body = next(l for l in prs.slide_layouts if l.name == "Title and Content")
        widths = [ph.width.inches for ph in body.placeholders if ph.placeholder_format.idx in (0, 1)]
        assert widths and all(abs(w - 11.93) < 0.05 for w in widths)


def test_pptx_chart_table_image_two_column_and_text_slides():
    with proj() as t:
        png(t / "logo.png", (800, 400))
        p = str(t / "d.pptx")
        out = call("pptx_create", path=p, slides=[
            {"title": "Revenue", "chart": {"type": "bar", "categories": ["Q1", "Q2"], "series": [{"name": "2025", "values": [10, "12"]}]}},
            {"title": "Share", "chart": {"type": "pie", "categories": ["A", "B"], "series": {"s": [60, 40]}}},
            {"title": "Compare", "left": ["- fast", "cheap"], "right": ["slow"]},
            {"title": "Logo", "image": (t / "logo.png").as_posix(), "caption": "Our logo", "bullets": ["Blue", "Simple"]},
            {"title": "Numbers", "table": [["a", "b"], ["1", "2"]], "notes": "say it slowly"},
            {"title": "Plain", "text": "One paragraph, no bullets."}])
        assert out.startswith("MEDIA:"), out
        txt = call("office_read", path=p)
        assert "[chart column_clustered: categories Q1, Q2; 2025: 10, 12]" in txt and "[chart pie: categories A, B; s: 60, 40]" in txt
        assert "- fast" in txt and "- cheap" in txt and "[picture" in txt and "Our logo" in txt and "| a | b |" in txt and "Notes: say it slowly" in txt
        assert "One paragraph, no bullets." in txt


def test_pptx_errors_are_specific():
    with proj() as t:
        p = str(t / "d.pptx")
        assert call("pptx_create", path=p, slides=[{"title": "x", "chart": {"type": "bar", "categories": ["a", "b", "c"], "series": [{"name": "s", "values": [1, 2]}]}}]) \
            .startswith("Error: slide 1: series 1 has 2 values but there are 3 categories")
        assert call("pptx_create", path=p, slides=[{"title": "x", "chart": {"type": "radar", "categories": ["a"], "series": [{"values": [1]}]}}]).startswith("Error: slide 1: chart type 'radar'")
        assert call("pptx_create", path=p, slides=[{"image": "x.png"}]).startswith("Error: slide 1: a 'image' slide needs a 'title'")
        assert call("pptx_create", path=p, slides=[{"title": "x", "image": "nope.png"}]).startswith("Error: slide 1: image not found")
        assert call("pptx_create", path=p).startswith("Error: nothing to write")
        assert call("pptx_create", path=p, markdown="   ").startswith("Error: nothing to write")
        assert not (t / "d.pptx").exists()


def test_pptx_edit_add_move_delete_replace_and_set_text():
    with proj() as t:
        p = str(t / "d.pptx"); call("pptx_create", path=p, markdown=DECK)
        out = call("pptx_edit", path=p, ops=[
            {"op": "add_slide", "slide": {"title": "Appendix", "bullets": ["extra"]}, "position": 2},
            {"op": "replace_text", "find": "Revenue up", "replace": "Revenue rose"},
            {"op": "move_slide", "slide": 2, "to": 5},
            {"op": "delete_slide", "slide": 4},
            {"op": "set_text", "slide": 2, "bullets": ["new one", "new two"], "notes": "changed"},
            {"op": "set_text", "slide": 1, "title": "Q3 Review (final)"}])
        assert out.startswith("MEDIA:"), out
        txt = call("office_read", path=p)
        titles = [l[7:] for l in txt.splitlines() if l.startswith("Title:")]
        assert titles == ["Q3 Review (final)", "Highlights", "Regions", "Appendix"]                      # Appendix went 2 -> 5, "Thank you" was deleted
        assert "- new one" in txt and "Notes: changed" in txt and "Revenue rose" not in txt              # set_text replaced that body
        import pptx
        assert len(pptx.Presentation(p).slides) == 4


def test_pptx_replace_text_reaches_runs_tables_and_notes_and_failures_change_nothing():
    with proj() as t:
        p = str(t / "d.pptx"); call("pptx_create", path=p, markdown=DECK)
        assert call("pptx_edit", path=p, ops=[{"op": "replace_text", "find": "Acme", "replace": "Globex"}, {"op": "replace_text", "find": "$900", "replace": "$950"}]).startswith("MEDIA:")
        txt = call("office_read", path=p)
        assert "Globex" in txt and "Acme" not in txt and "Mention Globex." in txt and "$950" in txt
        before = sha(p)
        assert call("pptx_edit", path=p, ops=[{"op": "delete_slide", "slide": 1}, {"op": "delete_slide", "slide": 99}]).startswith("Error: op 2: the deck has 3 slide(s); there is no slide 99")
        assert sha(p) == before
        assert call("pptx_edit", path=p, ops=[{"op": "set_text", "slide": 3}]).startswith("Error: op 1: set_text needs at least one of")
        assert call("pptx_edit", path=p, ops=[{"op": "set_text", "slide": 3, "bullets": ["x"]}]).startswith("Error: op 1: slide 3 has no text box")  # table slide


def test_pptx_edit_keeps_the_theme_of_decks_made_elsewhere():
    with proj() as t:
        import pptx
        src = pptx.Presentation()                                                        # a plain 4:3 deck from "somewhere else"
        s = src.slides.add_slide(src.slide_layouts[1]); s.shapes.title.text = "Old"; s.placeholders[1].text = "old text"
        p = str(t / "other.pptx"); src.save(p)
        assert call("pptx_edit", path=p, ops=[{"op": "add_slide", "slide": {"title": "New", "bullets": ["one", "two"]}}]).startswith("MEDIA:")
        prs = pptx.Presentation(p)
        assert len(prs.slides) == 2 and abs(prs.slide_width.inches - 10.0) < 0.01       # still 4:3, not forced to 16:9
        new = prs.slides[1]
        run = new.shapes.title.text_frame.paragraphs[0].runs[0]
        assert run.font.color.type is None and run.font.name is None                     # their fonts / colors, not ours
        assert not any(sh.name == "Slide Number" or sh.name.startswith("Rectangle") for sh in new.shapes)


# ------------------------------------------------------------------------------------------------ PDF
def test_pdf_create_read_back_pages_unicode_and_paper_size():
    with proj() as t:
        p = str(t / "p.pdf")
        md = "# PDF Test\n\nPrice is \u20b11,200 and caf\u00e9.\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n- x\n- y\n\n[pagebreak]\n\n## Page two\nSecond page text.\n"
        out = call("pdf_create", path=p, markdown=md)
        uni = pdf_tools._find_fonts()["unicode"]
        if not uni:                                                                      # no TrueType font on this machine: honest report
            assert "cannot be drawn" in out and "\u20b1" in out
            return
        assert out.startswith("MEDIA:"), out
        txt = call("office_read", path=p)
        assert "PDF: 2 page(s)" in txt and "--- Page 1 ---" in txt and "PDF Test" in txt and "\u20b11,200" in txt and "caf\u00e9" in txt and "Second page text." in txt
        assert "Second page" not in call("office_read", path=p, pages="1") and "Second page" in call("office_read", path=p, pages="2")
        call("pdf_create", path=str(t / "a4.pdf"), markdown="# x", paper="a4")
        import pypdf
        box = pypdf.PdfReader(str(t / "a4.pdf")).pages[0].mediabox
        assert abs(float(box.width) - 595) < 3 and abs(float(box.height) - 842) < 3


def test_pdf_with_blocks_image_and_numbers_and_bad_input():
    with proj() as t:
        png(t / "a.png")
        assert call("pdf_create", path=str(t / "p.pdf"), blocks=[{"type": "title", "text": "T"}, {"type": "numbered", "items": ["a", "b"]},
                                                              {"type": "image", "path": (t / "a.png").as_posix(), "caption": "Pic"}]).startswith(("MEDIA:", "Saved"))
        txt = call("office_read", path=str(t / "p.pdf"))
        assert "1." in txt and "2." in txt and "Pic" in txt
        assert call("pdf_create", path=str(t / "q.pdf")).startswith("Error: nothing to write")
        assert call("pdf_create", path=str(t / "q.docx"), markdown="# x").startswith("Error: a PDF file must end with .pdf")


# ------------------------------------------------------------------------------------------------ office_read
def test_office_read_csv_long_files_wrong_types_and_damaged_files():
    with proj() as t:
        (t / "d.csv").write_bytes("name,city\nAna,Cebu\nJos\u00e9,Manila\n".encode("utf-8-sig"))
        txt = call("office_read", path=str(t / "d.csv"))
        assert "CSV: 3 row(s) x 2 column(s)" in txt and "| Jos\u00e9 | Manila |" in txt
        call("docx_create", path=str(t / "long.docx"), markdown="\n\n".join(f"Paragraph {i} " + "word " * 40 for i in range(120)))
        first = call("office_read", path=str(t / "long.docx"))
        assert "[Cut here:" in first and "offset=12000" in first and "Paragraph 0" in first
        second = call("office_read", path=str(t / "long.docx"), offset=12000)
        assert second.startswith("[continuing from character 12000]") and "Paragraph 0 " not in second
        assert call("office_read", path=str(t / "long.docx"), offset=10_000_000).startswith("[nothing more")
        (t / "n.txt").write_text("hi")
        assert call("office_read", path=str(t / "n.txt")).startswith("Error: an Office / PDF / CSV file must end with")
        (t / "bad.docx").write_text("not a zip at all")
        assert "not a valid Office file" in call("office_read", path=str(t / "bad.docx"))
        assert call("office_read", path=str(t / "nope.docx")).startswith("Error: file not found")


def test_files_from_the_downloads_folder_are_read_as_untrusted_data():
    env = Env("ask_for_risky")
    saved = (_paths.SAFE_ROOT, _paths.ROOT, _paths.RUNTIME)
    _paths.SAFE_ROOT = _paths.ROOT = env.root; _paths.RUNTIME = env.root / "runtime"
    try:
        call("docx_create", path=env.p("work/ok.docx"), markdown="# Mine\n\nHello.")
        assert not call("office_read", path=env.p("work/ok.docx")).startswith("[UNTRUSTED")
        shutil.copy(env.p("work/ok.docx"), env.p("Downloads/from_a_stranger.docx"))
        out = call("office_read", path=env.p("Downloads/from_a_stranger.docx"))
        assert "# Mine" in unwrap(out)                                                    # framed as DATA, never as instructions
    finally:
        _paths.SAFE_ROOT, _paths.ROOT, _paths.RUNTIME = saved
        env.close()


# ------------------------------------------------------------------------------------------------ safety layer
def test_safety_rules_apply_to_the_target_file_and_to_pictures_named_inside_the_arguments():
    env = Env("ask_for_risky")
    saved = (_paths.SAFE_ROOT, _paths.ROOT, _paths.RUNTIME)
    _paths.SAFE_ROOT = _paths.ROOT = env.root; _paths.RUNTIME = env.root / "runtime"
    try:
        secret = png(env.root / "secrets" / "id.png")
        fine = png(env.root / "work" / "ok.png")
        # a picture from a blocked folder cannot be embedded (the policy cannot see it inside `blocks`, so the tool asks)
        out = call("docx_create", path=env.p("work/a.docx"), markdown=f"![x]({secret.as_posix()})")
        assert out.startswith("Error: block 1: cannot use id.png") and "blocked" in out and not (env.root / "work" / "a.docx").exists()
        assert call("pptx_create", path=env.p("work/a.pptx"), slides=[{"title": "x", "image": secret.as_posix()}]).startswith("Error: slide 1: cannot use id.png")
        assert call("docx_create", path=env.p("work/b.docx"), markdown=f"![x]({fine.as_posix()})").startswith("MEDIA:")
        # the target itself goes through the normal hard rules
        for name, args in [("docx_create", {"markdown": "# a"}), ("xlsx_create", {"sheets": [{"name": "s", "columns": ["a"], "rows": [[1]]}]}),
                           ("pptx_create", {"markdown": "# a"}), ("pdf_create", {"markdown": "# a"})]:
            ext = {"docx_create": "docx", "xlsx_create": "xlsx", "pptx_create": "pptx", "pdf_create": "pdf"}[name]
            res = call(name, path=env.p(f"secrets/x.{ext}"), **args)
            assert res.startswith("Refused by Lucy's safety rules") and not (env.root / "secrets" / f"x.{ext}").exists(), name
        assert call("office_read", path=env.p("secrets/x.docx")).startswith("Refused by Lucy's safety rules")
        # overwriting a file is undoable: the safety layer snapshots it first
        call("docx_create", path=env.p("work/b.docx"), markdown="# replaced")
        assert env.layer.backups.list(), "no undo snapshot was taken"
    finally:
        _paths.SAFE_ROOT, _paths.ROOT, _paths.RUNTIME = saved
        env.close()


# ------------------------------------------------------------------------------------------------ office_render
def test_office_render_draws_pdf_pages_without_libreoffice_and_gives_an_overview():
    with proj() as t:
        call("pdf_create", path=str(t / "p.pdf"), markdown="# One\n\ntext\n\n[pagebreak]\n\n# Two\n\n[pagebreak]\n\n# Three\n")
        out = call("office_render", path=str(t / "p.pdf"))
        assert out.startswith("Drew 3 of 3 page(s) of p.pdf") and "analyze_image" in out
        paths = [l.split(": ", 1)[1].split(" (")[0] for l in out.splitlines() if l.startswith("- ")]
        assert len(paths) == 4 and paths[-1].endswith("overview.png")                    # 3 pages + the overview
        from PIL import Image
        for p in paths:
            assert pathlib.Path(p).is_file() and Image.open(p).width > 100
        one = call("office_render", path=str(t / "p.pdf"), pages="2")
        assert one.startswith("Drew 1 of 3") and "page 2:" in one and "overview" not in one
        assert call("office_render", path=str(t / "p.pdf"), pages="9").startswith("Error: no such page: the file has 3 page(s)")
        assert call("office_render", path=str(t / "p.pdf"), pages="x-y").startswith("Error: cannot read the page list")


def test_office_render_asks_for_libreoffice_for_office_files_when_it_is_missing():
    with proj() as t:
        call("docx_create", path=str(t / "r.docx"), markdown="# Hi")
        real = O.find_soffice
        O.find_soffice = lambda: None
        try:
            out = call("office_render", path=str(t / "r.docx"))
        finally:
            O.find_soffice = real
        assert out.startswith("Error: LibreOffice is not installed") and "office_read" in out


def test_office_render_draws_word_and_powerpoint_files_with_libreoffice():
    if not HAVE_LO:
        return
    with proj() as t:
        call("docx_create", path=str(t / "r.docx"), markdown=MD)
        call("pptx_create", path=str(t / "d.pptx"), markdown=DECK)
        out = call("office_render", path=str(t / "d.pptx"), pages="1-2")
        assert out.startswith("Drew 2 of 4 page(s) of d.pptx") and "more pages" in out and "pages='3-4'" in out
        out = call("office_render", path=str(t / "r.docx"))
        assert out.startswith("Drew 2 of 2 page(s) of r.docx")
        again = call("office_render", path=str(t / "r.docx"))                             # second call reuses the converted PDF
        assert again.splitlines()[1] == out.splitlines()[1]


# ------------------------------------------------------------------------------------------------ routing + guidance
def test_routing_shows_only_the_office_tools_the_request_needs():
    def extra(text):
        return [n for n in routing.select_tools(text) if n not in routing.ALWAYS]
    assert {"docx_create", "office_render"} <= set(extra("make me a word document about our Q3 sales")) and "docx_edit" not in extra("make me a word document about sales")
    assert "xlsx_edit" in extra("update budget.xlsx and add a total column") and "xlsx_create" not in extra("update budget.xlsx and add a total column")
    assert extra("what's in report.docx?").count("office_read") == 1 and "docx_edit" not in extra("what's in report.docx?")
    assert {"pptx_create", "office_render"} <= set(extra("create a powerpoint deck about cats"))
    assert {"pptx_edit"} <= set(extra("change the title on slide 3"))
    assert "pdf_create" in extra("turn this into a PDF") and "pdf_create" not in extra("summarize contract.pdf")
    for chat in ["hello how are you", "write a poem about the sea", "what is 2+2"]:
        assert not set(extra(chat)) & routing.OFFICE_TOOLS and "office_read" not in extra(chat), chat
    assert len(routing.select_tools("make a word document about sales")) <= 11                # the prompt stays small


def test_tool_search_finds_the_office_tools():
    cat = routing.ToolCatalog(reg.schemas())
    names = lambda q: [t["function"]["name"] for t in cat.search(q)]
    assert "xlsx_create" in names("make an excel spreadsheet")
    assert "pptx_create" in names("create a presentation with slides")
    assert "office_read" in names("read a docx file")


def test_office_guidance_is_loaded_from_the_skill_file_without_its_front_matter_and_hot_reloads():
    with proj() as t:
        routing._guide_cache.update(mtime=None, text=routing._OFFICE_DEFAULT)
        assert routing.office_guidance() == routing._OFFICE_DEFAULT                       # no file: built-in text
        f = t / ".core" / "skills" / "definitions" / "office" / "SKILL.md"; f.parent.mkdir(parents=True)
        f.write_text("---\nname: office\ndescription: x\n---\nMy own office rules.\n", encoding="utf-8")
        assert routing.office_guidance() == "My own office rules."
        f.write_text("---\nname: office\n---\nChanged rules.\n", encoding="utf-8"); os.utime(f, (f.stat().st_atime, f.stat().st_mtime + 5))
        assert routing.office_guidance() == "Changed rules."
        routing._guide_cache.update(mtime=None, text=routing._OFFICE_DEFAULT)
        f.unlink(); old = t / ".core" / "skills" / "office" / "SKILL.md"; old.parent.mkdir(parents=True)
        old.write_text("Old place rules.", encoding="utf-8")                             # the earlier location still works
        assert routing.office_guidance() == "Old place rules."
        routing._guide_cache.update(mtime=None, text=routing._OFFICE_DEFAULT)
    real = (pathlib.Path(__file__).resolve().parents[1] / ".core" / "skills" / "definitions" / "office" / "SKILL.md").read_text(encoding="utf-8")
    assert real.startswith("---") and "office_render" in real and len(real) < 2600 + 400      # the shipped file fits the injection cap
    assert routing.office_mode(["docx_create"]) and routing.office_mode(["office_render"]) and not routing.office_mode(["office_read", "read_local_file"])


# ------------------------------------------------------------------------------------------------ Word strictness
PPT_PPR_ORDER = ["lnSpc", "spcBef", "spcAft", "buClrTx", "buClr", "buSzTx", "buSzPct", "buSzPts", "buFontTx", "buFont", "buNone", "buAutoNum",
                 "buChar", "buBlip", "tabLst", "defRPr", "extLst"]


def _order_violations(path):
    """Property elements out of schema order. LibreOffice does not care, WORD DOES: it answers 'unreadable content' (this once shipped)."""
    import zipfile
    from lxml import etree
    from lucy.tools.docx_tools import SCHEMA_ORDER, _local
    bad = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if not n.endswith(".xml"):
                continue
            for el in etree.fromstring(z.read(n)).iter():
                if not isinstance(el.tag, str):
                    continue
                ln = _local(el.tag)
                order = SCHEMA_ORDER.get(ln) if n.startswith("word/") else (PPT_PPR_ORDER if ln == "pPr" and n.startswith("ppt/") else None)
                if order:
                    idx = [order.index(_local(c.tag)) for c in el if isinstance(c.tag, str) and _local(c.tag) in order]
                    if idx != sorted(idx):
                        bad.append((n, ln, [_local(c.tag) for c in el]))
    return bad


def test_generated_word_and_powerpoint_files_follow_the_schema_order_that_word_demands():
    with proj() as t:
        png(t / "a.png")
        call("docx_create", path=str(t / "r.docx"), markdown=MD + f"\n\n![cap]({(t / 'a.png').as_posix()})")
        call("docx_create", path=str(t / "cv.docx"), style="resume", markdown="# Jane\n\ncontact\n\n## Experience\n- a\n\n1. x\n2. y\n")
        call("docx_create", path=str(t / "e.docx"), markdown=MD)
        assert call("docx_edit", path=str(t / "e.docx"), ops=[{"op": "add_row", "table": 1, "cells": ["a", "1"]}, {"op": "insert_after", "match": "Numbers", "markdown": "- a\n- b"},
                                                         {"op": "append", "markdown": "1. one\n2. two\n\n| a | b |\n|---|---|\n| 1 | 2 |"}, {"op": "set_cell", "table": 1, "row": 2, "col": 1, "text": "Z"}]).startswith("MEDIA:")
        call("pptx_create", path=str(t / "d.pptx"), markdown=DECK)
        call("pptx_create", path=str(t / "d2.pptx"), slides=[{"title": "c", "chart": {"type": "bar", "categories": ["a"], "series": [{"name": "s", "values": [1]}]}},
                                                          {"title": "i", "image": (t / "a.png").as_posix(), "bullets": ["x"]}, {"title": "two", "left": ["a"], "right": ["b"]},
                                                          {"title": "plain", "text": "no bullets"}])
        call("pptx_edit", path=str(t / "d.pptx"), ops=[{"op": "add_slide", "slide": {"title": "x", "text": "plain"}}, {"op": "move_slide", "slide": 1, "to": 3}])
        for f in ("r.docx", "cv.docx", "e.docx", "d.pptx", "d2.pptx"):
            assert _order_violations(t / f) == [], f
        import docx
        d = docx.Document(str(t / "r.docx"))
        z = d.settings.element.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}zoom")
        assert z is not None and z.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}percent") == "100"


def _pptx_package_problems(path):
    """What PowerPoint repairs but the schema validator cannot see: duplicate part names, dangling links, slide count mismatch."""
    import collections, posixpath, re, zipfile
    from lxml import etree
    z = zipfile.ZipFile(path)
    names = z.namelist()
    probs = [f"duplicate part {k}" for k, v in collections.Counter(names).items() if v > 1]
    have = set(names)
    for n in have:
        if n.endswith(".rels"):
            base = posixpath.dirname(posixpath.dirname(n))
            for r in etree.fromstring(z.read(n)):
                if r.get("TargetMode") == "External":
                    continue
                tgt = r.get("Target")
                full = posixpath.normpath(tgt.lstrip("/")) if tgt.startswith("/") else posixpath.normpath(posixpath.join(base, tgt))
                if full not in have:
                    probs.append(f"dangling link {n} -> {full}")
    listed = len(re.findall(rb"<p:sldId ", z.read("ppt/presentation.xml")))
    parts = len([n for n in have if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)])
    if listed != parts:
        probs.append(f"{listed} slides listed but {parts} slide parts")
    return probs


def test_pptx_edits_that_delete_and_add_never_leave_two_parts_with_the_same_name():
    with proj() as t:
        png(t / "a.png")
        for deck, ops in [
            ("# A\n\n---\n## B\n- x\n\n---\n## C\n- y\n", [{"op": "delete_slide", "slide": 2}, {"op": "add_slide", "slide": {"title": "n", "bullets": ["z"]}}]),
            ("# A\n\n---\n## B\n- x\nNotes: n1\n\n---\n## C\n- y\n", [{"op": "delete_slide", "slide": 2}]),
            (DECK, [{"op": "add_slide", "slide": {"title": "n"}}, {"op": "delete_slide", "slide": 1}, {"op": "add_slide", "slide": {"title": "m", "image": (t / "a.png").as_posix()}},
                    {"op": "move_slide", "slide": 1, "to": 3}, {"op": "delete_slide", "slide": 2}, {"op": "add_slide", "slide": {"title": "k", "chart": {"type": "pie", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}}}]),
        ]:
            p = str(t / "d.pptx")
            call("pptx_create", path=p, markdown=deck)
            assert call("pptx_edit", path=p, ops=ops).startswith("MEDIA:")
            assert _pptx_package_problems(p) == [], ops
            assert _order_violations(p) == []
        assert _pptx_package_problems(p) == []
