"""office_read and office_render: SEE what is in an Office / PDF file.

office_read returns readable text for docx / xlsx / pptx / pdf / csv (old .doc .xls .ppt .odt .rtf go through LibreOffice when it is
installed): headings, bullets, tables (as | pipe | rows), slide by slide, sheet by sheet, page by page. Long files are cut at
~12k characters with a note saying how to continue. office_render turns pages into PNG pictures (LibreOffice -> PDF -> PNG) so the
model can LOOK at its own work with analyze_image (visual QA loop): text cut off? overlapping? ugly table?
"""
from __future__ import annotations

import csv
import datetime as _dt
import os
import re
import shutil
import tempfile
import time
import zipfile
from pathlib import Path
from typing import Any, List, Optional, Tuple

from lucy import paths as _paths

from . import _office as O
from .registry import ToolContext

CAP = 12000
READ_EXT = (".docx", ".xlsx", ".xlsm", ".pptx", ".pdf", ".csv", ".tsv", ".doc", ".xls", ".ppt", ".odt", ".ods", ".odp", ".rtf")
LEGACY = {".doc": "docx", ".odt": "docx", ".rtf": "docx", ".xls": "xlsx", ".ods": "xlsx", ".ppt": "pptx", ".odp": "pptx"}
RENDER_EXT = READ_EXT[:-1] + (".rtf",)

SCHEMA_READ = {
    "type": "function",
    "function": {
        "name": "office_read",
        "description": ("Read a Word, Excel, PowerPoint, PDF or CSV file as text (headings, bullets, tables, slides, sheets, pages). "
                        "Use it before editing a file and to read files the user gives you."),
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "The file to read."},
                "pages": {"type": "string", "description": "PDF pages or slides to read, e.g. 2 or 1-3 (default: all, up to the size limit)."},
                "sheet": {"type": "string", "description": "Excel: the sheet name (default: every sheet)."},
                "range": {"type": "string", "description": "Excel: the cells to read, e.g. A1:F40 (default: the first 50 rows)."},
                "offset": {"type": "integer", "description": "To continue a long answer: the character number where the last one stopped."},
            },
            "required": ["path"],
        },
    },
}

SCHEMA_RENDER = {
    "type": "function",
    "function": {
        "name": "office_render",
        "description": ("Turn pages / slides of a document into PNG pictures so you can check the layout with analyze_image "
                        "(text cut off? overlapping? ugly?). Needs LibreOffice for Word/Excel/PowerPoint; PDFs work without."),
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "The file to draw."},
                "pages": {"type": "string", "description": "Which pages / slides, e.g. 1-3 (default 1-4, at most 8 at a time)."},
            },
            "required": ["path"],
        },
    },
}


# ------------------------------------------------------------------------------------------- text helpers
def _clip(s: Any, n: int = 80) -> str:
    t = re.sub(r"\s+", " ", str(s)).strip()
    return t if len(t) <= n else t[:n - 1] + "\u2026"


def _num(v: float) -> str:
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int) or float(v).is_integer():
        return str(int(v))
    return f"{v:.6g}"


def _cell_text(v: Any, fmt: str = "") -> str:
    if v is None:
        return ""
    if isinstance(v, _dt.datetime):
        return v.date().isoformat() if v.time() == _dt.time(0) else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, _dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if "%" in (fmt or ""):
            return f"{v * 100:.1f}".rstrip("0").rstrip(".") + "%"
        return _num(v)
    return _clip(v, 80)


def _table_md(rows: List[List[str]]) -> List[str]:
    out = []
    for i, r in enumerate(rows):
        out.append("| " + " | ".join(c.replace("|", "/").replace("\n", " ") for c in r) + " |")
        if i == 0 and len(rows) > 1:
            out.append("|" + "|".join("---" for _ in r) + "|")
    return out


def _window(text: str, offset: int) -> str:
    offset = max(0, int(offset or 0))
    if offset >= len(text) > 0:
        return f"[nothing more: the text is {len(text)} characters long]"
    chunk = text[offset:offset + CAP]
    if len(text) > offset + CAP:
        chunk += f"\n\n[Cut here: {len(text) - offset - CAP} more characters. Call office_read again with offset={offset + CAP} to continue.]"
    if offset:
        chunk = f"[continuing from character {offset}]\n" + chunk
    return chunk


# ------------------------------------------------------------------------------------------- readers
def _read_docx(path: Path) -> str:
    docx = O.need("docx", "python-docx")
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn
    doc = docx.Document(str(path))
    out: List[str] = []
    for part, label in ((doc.sections[0].header, "header"), (doc.sections[0].footer, "footer")) if doc.sections else ():
        t = " ".join(p.text.strip() for p in part.paragraphs if p.text.strip())
        if t:
            out.append(f"[{label}] {_clip(t, 200)}")
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            text = p.text.strip()
            style = (p.style.name if p.style is not None else "") or ""
            has_img = bool(el.findall(".//" + qn("w:drawing"))) or bool(el.findall(".//" + qn("w:pict")))
            has_break = any(br.get(qn("w:type")) == "page" for br in el.findall(".//" + qn("w:br")))
            if has_break and not text:
                out.append("[page break]")
                continue
            if not text and not has_img:
                continue
            low = style.lower()
            m = re.match(r"heading\s*(\d)", low)
            if low == "title":
                line = "# " + text
            elif m:
                line = "#" * min(6, int(m.group(1)) + 1) + " " + text
            elif low.startswith("list bullet") or low.startswith("list paragraph"):
                lvl = int(low[-1]) - 1 if low[-1].isdigit() else 0
                line = "  " * lvl + "- " + text
            elif low.startswith("list number"):
                lvl = int(low[-1]) - 1 if low[-1].isdigit() else 0
                line = "  " * lvl + "1. " + text
            elif el.find(".//" + qn("w:numPr")) is not None:
                ilvl = el.find(".//" + qn("w:ilvl"))
                lvl = int(ilvl.get(qn("w:val"))) if ilvl is not None else 0
                line = "  " * lvl + "- " + text
            else:
                line = text
            out.append(line + (" [image]" if has_img else ""))
        elif el.tag == qn("w:tbl"):
            t = Table(el, doc)
            rows = []
            for r in t.rows:
                seen, cells = set(), []
                for c in r.cells:
                    if id(c._tc) in seen:
                        continue
                    seen.add(id(c._tc))
                    cells.append(_clip(c.text, 120))
                rows.append(cells)
            if rows:
                w = max(len(r) for r in rows)
                out.append("")
                out += _table_md([r + [""] * (w - len(r)) for r in rows])
                out.append("")
    return "\n".join(out).strip() or "(the document has no text)"


def _read_xlsx(path: Path, sheet: Optional[str], rng: Optional[str]) -> str:
    openpyxl = O.need("openpyxl", "openpyxl")
    from openpyxl.utils import get_column_letter, range_boundaries
    wb = openpyxl.load_workbook(path)
    vals = openpyxl.load_workbook(path, data_only=True)
    has_formula = any(isinstance(c.value, str) and c.value.startswith("=") for ws in wb.worksheets for row in ws.iter_rows() for c in row)
    stale = has_formula and any(isinstance(c.value, str) and c.value.startswith("=") and vals[ws.title][c.coordinate].value is None
                                for ws in wb.worksheets for row in ws.iter_rows() for c in row)
    note = ""
    if stale:
        fresh = O.recalculated(path)
        if fresh is not None:
            vals = fresh
        else:
            note = "(formula results are not stored in this file and LibreOffice is not available to calculate them, so only formulas are shown)"
    names = [ws.title for ws in wb.worksheets]
    if sheet:
        pick = [n for n in names if n.lower() == str(sheet).strip().lower()]
        if not pick:
            raise O.OfficeError(f"no sheet '{sheet}'. The sheets are: {', '.join(names)}.")
        names = pick
    out: List[str] = []
    try:
        with zipfile.ZipFile(path) as z:
            charts = sum(1 for n in z.namelist() if n.startswith("xl/charts/chart"))
    except Exception:                                                # noqa: BLE001
        charts = 0
    out.append(f"Workbook: {len(wb.worksheets)} sheet(s): {', '.join(w.title for w in wb.worksheets)}" + (f"; {charts} chart(s)" if charts else ""))
    if note:
        out.append(note)
    for name in names:
        ws, vs = wb[name], vals[name] if name in vals.sheetnames else None
        dims = ws.dimensions if ws.max_row > 1 or ws.max_column > 1 or ws["A1"].value is not None else "(empty)"
        out.append("")
        out.append(f"## Sheet \"{name}\" {dims}" + (" (hidden)" if ws.sheet_state != "visible" else "") +
                   (f" freeze={ws.freeze_panes}" if ws.freeze_panes else ""))
        if dims == "(empty)":
            continue
        if rng:
            try:
                c1, r1, c2, r2 = range_boundaries(str(rng).upper().replace("$", ""))
            except ValueError:
                raise O.OfficeError(f"cannot read the range '{rng}'. Use e.g. A1:F40.")
            more = None
        else:
            c1, r1, c2, r2 = 1, 1, min(ws.max_column, 30), min(ws.max_row, 50)
            more = ws.max_row - r2 if ws.max_row > r2 else None
        c2 = min(c2, c1 + 29)
        r2 = min(r2, r1 + 199)
        out.append("    | " + " | ".join(get_column_letter(c) for c in range(c1, c2 + 1)) + " |")
        for r in range(r1, r2 + 1):
            cells = []
            for c in range(c1, c2 + 1):
                cell = ws.cell(row=r, column=c)
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    got = vs[cell.coordinate].value if vs is not None else None
                    shown = _cell_text(got, cell.number_format)
                    cells.append(f"{_clip(v, 50)} \u2192 {shown}" if got is not None else _clip(v, 60))
                else:
                    cells.append(_cell_text(v, cell.number_format))
            if any(cells):
                out.append(f"{r:>3} | " + " | ".join(cells) + " |")
        if more:
            out.append(f"... {more} more row(s). Call again with range='A{r2 + 1}:{get_column_letter(min(ws.max_column, 30))}{min(ws.max_row, r2 + 50)}'.")
    return "\n".join(out)


def _read_pptx(path: Path, pages: Optional[str]) -> str:
    pptx = O.need("pptx", "python-pptx")
    prs = pptx.Presentation(str(path))
    total = len(prs.slides)
    picks = O.page_range(pages, total, max_pages=60)
    out = [f"Deck: {total} slide(s)" + (f", showing {picks[0]}-{picks[-1]}" if len(picks) < total else "")]

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                yield from walk(sh.shapes)
            else:
                yield sh
    for n in picks:
        slide = prs.slides[n - 1]
        out.append("")
        out.append(f"--- Slide {n} ({slide.slide_layout.name}) ---")
        title = slide.shapes.title
        if title is not None and title.has_text_frame and title.text_frame.text.strip():
            out.append("Title: " + _clip(title.text_frame.text, 200))
        for sh in walk(slide.shapes):
            if title is not None and sh.shape_id == title.shape_id:
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                is_body = sh.is_placeholder and sh.placeholder_format.idx in (1, 2)
                for par in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in par.runs).strip()
                    if t and t != "\u2039#\u203a" and not (sh.name == "Slide Number"):
                        out.append("  " * par.level + ("- " if is_body else "") + _clip(t, 300))
            elif getattr(sh, "has_table", False) and sh.has_table:
                out += _table_md([[_clip(c.text, 80) for c in r.cells] for r in sh.table.rows])
            elif getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                try:
                    cats = [str(c) for c in ch.plots[0].categories]
                    ser = "; ".join(f"{s.name}: " + ", ".join(_num(v) for v in s.values) for s in ch.plots[0].series)
                    out.append(f"[chart {str(ch.chart_type).split('.')[-1].split(' ')[0].lower()}: categories {', '.join(cats)}; {ser}]")
                except Exception:                                    # noqa: BLE001
                    out.append("[chart]")
            elif sh.shape_type == 13:
                out.append(f"[picture {sh.name}]")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("Notes: " + _clip(slide.notes_slide.notes_text_frame.text, 500))
    return "\n".join(out)


def _read_pdf(path: Path, pages: Optional[str]) -> str:
    pypdf = O.need("pypdf", "pypdf")
    rd = pypdf.PdfReader(str(path))
    if rd.is_encrypted:
        try:
            if not rd.decrypt(""):
                raise O.OfficeError(f"{path.name} is password protected. Ask the user for the password or an unprotected copy.")
        except O.OfficeError:
            raise
        except Exception as e:                                       # noqa: BLE001
            raise O.OfficeError(f"{path.name} is encrypted and could not be opened ({type(e).__name__}).") from e
    total = len(rd.pages)
    picks = O.page_range(pages, total, max_pages=total)
    out = [f"PDF: {total} page(s)" + (f", showing {picks[0]}-{picks[-1]}" if len(picks) < total else "")]
    empty = 0
    for n in picks:
        try:
            t = (rd.pages[n - 1].extract_text() or "").strip()
        except Exception:                                            # noqa: BLE001
            t = ""
        if not t:
            empty += 1
        out.append("")
        out.append(f"--- Page {n} ---")
        out.append(t or "(no text on this page)")
    if empty == len(picks):
        out.append("")
        out.append("This PDF has no text layer (scanned images?). Use office_render to get page pictures, then analyze_image to read them.")
    return "\n".join(out)


def _read_csv(path: Path, rng: Optional[str]) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            txt = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    delim = "\t" if path.suffix.lower() == ".tsv" else (csv.Sniffer().sniff(txt[:4000], delimiters=",;\t|").delimiter if txt.strip() else ",")
    rows = list(csv.reader(txt.splitlines(), delimiter=delim))
    out = [f"CSV: {len(rows)} row(s) x {max((len(r) for r in rows), default=0)} column(s)"]
    shown = rows[:60]
    w = max((len(r) for r in shown), default=0)
    out += _table_md([[_clip(c, 60) for c in r] + [""] * (w - len(r)) for r in shown])
    if len(rows) > 60:
        out.append(f"... {len(rows) - 60} more row(s). Use python or run_shell_command for the rest.")
    return "\n".join(out)


def _read(args: dict, ctx: ToolContext) -> str:
    tmpdir: Optional[Path] = None
    try:
        src = O.target_path(args.get("path"), READ_EXT, must_exist=True, what="an Office / PDF / CSV file")
        path = src
        ext = src.suffix.lower()
        if ext in LEGACY:
            tmpdir = Path(tempfile.mkdtemp(prefix="lucy_conv_"))
            path = O.soffice_convert(src, LEGACY[ext], tmpdir)
            ext = path.suffix.lower()
        if ext == ".docx":
            text = _read_docx(path)
        elif ext in (".xlsx", ".xlsm"):
            text = _read_xlsx(path, args.get("sheet"), args.get("range"))
        elif ext == ".pptx":
            text = _read_pptx(path, args.get("pages"))
        elif ext == ".pdf":
            text = _read_pdf(path, args.get("pages"))
        else:
            text = _read_csv(path, args.get("range"))
        return _window(text, args.get("offset") or 0)
    except Exception as e:                                           # noqa: BLE001
        if not isinstance(e, O.OfficeError) and any(k in type(e).__name__.lower() for k in ("zip", "packagenotfound", "pdfread", "pdfstream")):
            return f"Error: {Path(str(args.get('path'))).name} is not a valid Office file (it may be damaged, or an old format renamed)."
        return O.fail(e)
    finally:
        if tmpdir:
            shutil.rmtree(tmpdir, ignore_errors=True)


# ------------------------------------------------------------------------------------------- office_render
def _safe(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", s)[:40] or "doc"


def _prune(root: Path, keep_hours: int = 24) -> None:
    if not root.is_dir():
        return
    cutoff = time.time() - keep_hours * 3600
    for d in root.iterdir():
        try:
            if d.is_dir() and d.stat().st_mtime < cutoff:
                shutil.rmtree(d, ignore_errors=True)
        except OSError:
            pass


def _overview(pngs: List[Path], out: Path) -> Optional[Path]:
    PIL = O.need("PIL.Image", "pillow")
    from PIL import ImageDraw
    try:
        thumbs = []
        for p in pngs:
            im = PIL.open(p).convert("RGB")
            h = int(im.height * 420 / im.width)
            thumbs.append((p, im.resize((420, h))))
        cols = min(3, len(thumbs))
        rows = (len(thumbs) + cols - 1) // cols
        th = max(t[1].height for t in thumbs)
        sheet = PIL.new("RGB", (cols * 430 + 10, rows * (th + 30) + 10), (120, 120, 120))
        for i, (p, im) in enumerate(thumbs):
            x, y = 10 + (i % cols) * 430, 10 + (i // cols) * (th + 30)
            sheet.paste(im, (x, y + 20))
            ImageDraw.Draw(sheet).text((x, y + 4), p.stem.replace("page-", "page "), fill=(255, 255, 255))
        sheet.save(out, "PNG")
        return out
    except Exception:                                                # noqa: BLE001
        return None


def _render(args: dict, ctx: ToolContext) -> str:
    try:
        src = O.target_path(args.get("path"), RENDER_EXT, must_exist=True, what="an Office / PDF file")
        root = _paths.RUNTIME / "tmp" / "render"
        _prune(root)
        out_dir = root / f"{_safe(src.stem)}-{O.file_key(src)}"
        out_dir.mkdir(parents=True, exist_ok=True)
        if src.suffix.lower() == ".pdf":
            pdf = src
        else:
            pdf = out_dir / (src.stem + ".pdf")
            if not pdf.is_file():
                made = O.soffice_convert(src, "pdf", out_dir)
                pdf = made
        total = O.pdf_page_count(pdf)
        pages = O.page_range(args.get("pages"), total, default=(1, 4), max_pages=8)
        try:
            dpi = max(50, min(150, int(os.environ.get("LUCY_RENDER_DPI", "80"))))
        except ValueError:
            dpi = 80
        pngs = O.render_pdf_pages(pdf, pages, out_dir, dpi=dpi)
        if not pngs:
            raise O.OfficeError("no page could be drawn.")
        lines = [f"Drew {len(pngs)} of {total} page(s) of {src.name} ({dpi} dpi). Check them with analyze_image, one picture per call "
                 "(ask e.g. 'Is any text cut off, overlapping or hard to read? Does the layout look tidy?'):"]
        for p in pngs:
            num = re.search(r"(\d+)", p.stem)
            lines.append(f"- page {num.group(1) if num else '?'}: {p.as_posix()}")
        if len(pngs) > 1:
            ov = _overview(pngs, out_dir / "overview.png")
            if ov:
                lines.append(f"- all pages on one picture: {ov.as_posix()} (good for a first quick look)")
        if total > pages[-1]:
            lines.append(f"There are more pages: call office_render again with pages='{pages[-1] + 1}-{min(total, pages[-1] + 4)}'.")
        return "\n".join(lines)
    except Exception as e:                                           # noqa: BLE001
        return O.fail(e)


TOOLS = [
    {"schema": SCHEMA_READ, "handler": _read, "tier": "read", "label": "read"},
    {"schema": SCHEMA_RENDER, "handler": _render, "tier": "read", "label": "analyze"},
]
