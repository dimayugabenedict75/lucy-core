"""Word documents: docx_create and docx_edit (python-docx, no Office needed).

The model never writes library code: it sends `blocks` (heading/para/bullets/numbered/table/image/pagebreak) or plain
`markdown`, picks a house `style` (report | resume), and the harness builds a clean file. Edits are a short list of ops,
applied all-or-nothing. A finished file is delivered to the chat (MEDIA:) so it also lands in the Library.
"""
from __future__ import annotations

import copy
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

from . import _office as O
from .registry import ToolContext

# ------------------------------------------------------------------------------------------- house templates
THEMES: Dict[str, Dict[str, Any]] = {
    # clean report: Calibri 11, blue headings, 1" margins, page numbers
    "report": dict(font="Calibri", size=11, text="222222", head="1F3864", h=(20, 15, 12), title=26, margin=1.0,
                   after=6, line=1.15, page_numbers=True, rule_h1=False, caps_h1=False, th_fill="1F3864", zebra="F3F6FA"),
    # one-page-friendly resume: tight margins, small text, ruled section headings
    "resume": dict(font="Calibri", size=10, text="222222", head="1F3864", h=(11.5, 11, 10), title=22, margin=0.7,
                   after=2, line=1.0, page_numbers=False, rule_h1=True, caps_h1=True, th_fill="1F3864", zebra="F3F6FA"),
}
_ALIASES = {"clean": "report", "clean report": "report", "default": "report", "cv": "resume", "plain": "report", "letter": "report"}
PAPERS = {"letter": (8.5, 11.0), "a4": (8.27, 11.69)}

SCHEMA_CREATE = {
    "type": "function",
    "function": {
        "name": "docx_create",
        "description": ("Create a Word document (.docx) from 'markdown' (easiest) or 'blocks'. The file is sent to the chat; "
                        "style 'report' or 'resume' gives it a clean layout."),
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Where to save it, e.g. C:/Users/you/Documents/report.docx"},
                "markdown": {"type": "string", "description": "The content: # Title, ## Heading, paragraphs, - bullets, 1. numbers, | tables |, ![caption](image path), [pagebreak]."},
                "blocks": {
                    "type": "array",
                    "description": "Instead of markdown: a list of blocks in order.",
                    "items": {
                        "type": "object",
                        "properties": {
                            "type": {"type": "string", "enum": ["title", "heading", "para", "bullets", "numbered", "table", "image", "pagebreak"]},
                            "text": {"type": "string", "description": "For title/heading/para. **bold** and *italic* work."},
                            "level": {"type": "integer", "description": "Heading level 1-3."},
                            "items": {"type": "array", "items": {"type": "string"}, "description": "For bullets/numbered."},
                            "rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}, "description": "For table: rows of cells, first row = header."},
                            "path": {"type": "string", "description": "For image: the picture file."},
                            "caption": {"type": "string", "description": "For image: caption text."},
                            "align": {"type": "string", "enum": ["left", "center", "right", "justify"]},
                        },
                        "required": ["type"],
                    },
                },
                "style": {"type": "string", "enum": ["report", "resume"], "description": "House layout (default report)."},
                "paper": {"type": "string", "enum": ["letter", "a4"], "description": "Paper size (default letter)."},
            },
            "required": ["path"],
        },
    },
}

SCHEMA_EDIT = {
    "type": "function",
    "function": {
        "name": "docx_edit",
        "description": ("Change an existing Word document with a list of 'ops' (all or none). Read it first with office_read so "
                        "your 'find' / 'match' text is exact."),
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "The .docx file (must exist)."},
                "ops": {
                    "type": "array",
                    "description": ("Changes in order. replace {find, replace}; append {markdown|blocks}; insert_after / insert_before "
                                    "{match, markdown|blocks}; delete {match}; set_cell {table, row, col, text}; add_row {table, cells}. "
                                    "Table/row/col numbers start at 1."),
                    "items": {
                        "type": "object",
                        "properties": {
                            "op": {"type": "string", "enum": ["replace", "append", "insert_after", "insert_before", "delete", "set_cell", "add_row"]},
                            "find": {"type": "string", "description": "replace: text to find."},
                            "replace": {"type": "string", "description": "replace: new text."},
                            "match": {"type": "string", "description": "insert/delete: a few words from the target paragraph."},
                            "markdown": {"type": "string", "description": "append/insert: new content as markdown."},
                            "blocks": {"type": "array", "items": {"type": "object"}, "description": "append/insert: new content as blocks."},
                            "table": {"type": "integer"}, "row": {"type": "integer"}, "col": {"type": "integer"},
                            "text": {"type": "string", "description": "set_cell: new text."},
                            "cells": {"type": "array", "items": {"type": "string"}, "description": "add_row: the new row's cells."},
                            "all": {"type": "boolean", "description": "replace/delete: every match (replace default true, delete default false)."},
                        },
                        "required": ["op"],
                    },
                },
            },
            "required": ["path", "ops"],
        },
    },
}


# ------------------------------------------------------------------------------------------- python-docx plumbing
def _libs():
    docx = O.need("docx", "python-docx")
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    return docx, WD_ALIGN_PARAGRAPH, OxmlElement, qn, Inches, Pt, RGBColor


# Child order the OOXML schema demands (Word REJECTS a file whose properties are out of order; LibreOffice does not care).
# python-docx offers no helper for elements it does not model, so the order is spelled out here.
SCHEMA_ORDER = {
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd",
            "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi",
            "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
            "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing",
              "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription", "tblPrChange"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign",
             "hideMark", "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange"],
    "trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc",
             "hidden", "ins", "del", "trPrChange"],
}


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _ordered_insert(parent, child) -> None:
    """Insert `child` into `parent` where the schema wants it: before the first existing child that must come AFTER it."""
    order = SCHEMA_ORDER.get(_local(parent.tag))
    mine = _local(child.tag)
    if order and mine in order:
        later = set(order[order.index(mine) + 1:])
        for existing in parent:
            if _local(existing.tag) in later:
                existing.addprevious(child)
                return
    parent.append(child)


def _hex(s: str):
    from docx.shared import RGBColor
    return RGBColor.from_string(s.upper())


def _style_font(style, name: str, size: Optional[float] = None, color: Optional[str] = None, bold: Optional[bool] = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is not None:
        for a in list(rfonts.attrib):
            if "theme" in a.lower():                        # theme fonts would override the name we set
                del rfonts.attrib[a]
        rfonts.set(qn("w:eastAsia"), name)
        rfonts.set(qn("w:cs"), name)
    if size:
        f.size = Pt(size)
    if color:
        f.color.rgb = _hex(color)
    if bold is not None:
        f.bold = bold


def _bottom_rule(style, color: str, sz: int = 6) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    ppr = style.element.get_or_add_pPr()
    for old in ppr.findall(qn("w:pBdr")):
        ppr.remove(old)
    b = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", str(sz)), ("space", "1"), ("color", color)):
        bt.set(qn(f"w:{k}"), v)
    b.append(bt)
    _ordered_insert_ppr(ppr, b)


def _ordered_insert_ppr(ppr, el) -> None:
    _ordered_insert(ppr, el)


def _no_border(style) -> None:
    from docx.oxml.ns import qn
    ppr = style.element.get_or_add_pPr()
    for old in ppr.findall(qn("w:pBdr")):
        ppr.remove(old)


def apply_theme(doc, theme: Dict[str, Any], paper: str) -> None:
    from docx.shared import Inches, Pt
    t = theme
    w, h = PAPERS.get(paper, PAPERS["letter"])
    for s in doc.sections:
        s.page_width, s.page_height = Inches(w), Inches(h)
        s.left_margin = s.right_margin = Inches(t["margin"])
        s.top_margin = s.bottom_margin = Inches(t["margin"] * (0.9 if t["margin"] < 1 else 1))
    st = doc.styles
    _style_font(st["Normal"], t["font"], t["size"], t["text"])
    st["Normal"].paragraph_format.space_after = Pt(t["after"])
    st["Normal"].paragraph_format.line_spacing = t["line"]
    for i, lvl in enumerate(("Heading 1", "Heading 2", "Heading 3")):
        s = st[lvl]
        _style_font(s, t["font"], t["h"][i], t["head"], True)
        s.paragraph_format.space_before = Pt(14 if i == 0 and not t["rule_h1"] else 9 if t["rule_h1"] else 10 - i * 2)
        s.paragraph_format.space_after = Pt(3 if t["rule_h1"] else 4)
        s.paragraph_format.keep_with_next = True
        s.font.italic = False
    if t["caps_h1"]:
        st["Heading 1"].font.all_caps = True
    if t["rule_h1"]:
        _bottom_rule(st["Heading 1"], t["head"], 6)
    ti = st["Title"]
    _style_font(ti, t["font"], t["title"], t["head"] if not t["rule_h1"] else "111111", True)
    ti.paragraph_format.space_after = Pt(4 if t["rule_h1"] else 10)
    if t["rule_h1"]:
        _no_border(ti)
    else:
        _bottom_rule(ti, t["head"], 8)
    sub = st["Subtitle"]
    _style_font(sub, t["font"], t["size"] + 2, "595959", False)
    sub.font.italic = False
    for name in ("List Bullet", "List Bullet 2", "List Bullet 3", "List Number", "List Number 2", "List Number 3"):
        try:
            s = st[name]
        except KeyError:
            continue
        _style_font(s, t["font"], t["size"], t["text"])
        s.paragraph_format.space_after = Pt(max(1, t["after"] - 2))
        s.paragraph_format.line_spacing = t["line"]


def add_runs(par, text: str, *, bold: bool = False, italic: bool = False, size: Optional[float] = None,
             color: Optional[str] = None) -> None:
    from docx.shared import Pt
    for seg, b, i, code in O.parse_inline(str(text)):
        if seg == "":
            continue
        r = par.add_run(seg)
        r.bold = True if (bold or b) else None
        r.italic = True if (italic or i) else None
        if code:
            r.font.name = "Consolas"
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = _hex(color)


def _new_restarted_num(doc, par) -> Optional[str]:
    """A fresh Word list (numId) that starts at 1, built on the same numbering definition as the paragraph's list style.
    Every level-0 item of one numbered block gets it, so the block counts 1, 2, 3 and the next block starts over."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    try:
        ppr = par.style.element.pPr
        npr = ppr.numPr if ppr is not None else None
        num_id = npr.numId.val if npr is not None and npr.numId is not None else None
        numbering = doc.part.numbering_part.element
        if num_id is None:
            return None
        abs_id = None
        for n in numbering.findall(qn("w:num")):
            if n.get(qn("w:numId")) == str(num_id):
                abs_id = n.find(qn("w:abstractNumId")).get(qn("w:val"))
        if abs_id is None:
            return None
        new_id = str(max([int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))] + [0]) + 1)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), new_id)
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), abs_id)
        num.append(a)
        ov = OxmlElement("w:lvlOverride")
        ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride")
        so.set(qn("w:val"), "1")
        ov.append(so)
        num.append(ov)
        numbering.append(num)
        return new_id
    except Exception:                                          # noqa: BLE001 - numbering still works, it just continues counting
        return None


def _use_num(par, num_id: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    ppr = par._p.get_or_add_pPr()
    for old in ppr.findall(qn("w:numPr")):
        ppr.remove(old)
    np_ = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl")
    il.set(qn("w:val"), "0")
    ni = OxmlElement("w:numId")
    ni.set(qn("w:val"), num_id)
    np_.append(il)
    np_.append(ni)
    _ordered_insert(ppr, np_)


_NUMBERISH = re.compile(r"^[\(\-+]?\s*[$\u20ac\u00a3\u20b1\u00a5]?\s*[\d,]+(\.\d+)?\s*%?\)?$")


def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _ordered_insert(tcpr, shd)


def _table_borders(table, color: str = "BFBFBF") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tblpr = table._tbl.tblPr
    for old in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(old)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", color)):
            e.set(qn(f"w:{k}"), v)
        b.append(e)
    _ordered_insert(tblpr, b)


def _repeat_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    _ordered_insert(trpr, h)


def _content_width(doc) -> float:
    s = doc.sections[-1]
    return (s.page_width - s.left_margin - s.right_margin) / 914400.0


def add_table(doc, rows: List[List[str]], header: bool, widths: Optional[list], theme: Dict[str, Any]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
    ncols = len(rows[0])
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.style = "Table Grid"
    tbl.autofit = False
    _table_borders(tbl)
    total = _content_width(doc)
    if widths and len(widths) == ncols:
        try:
            ws = [max(0.1, float(x)) for x in widths]
            ws = [w / sum(ws) * total for w in ws]
        except (TypeError, ValueError):
            ws = None
    else:
        ws = None
    if ws is None:
        lens = [max(4, min(40, max(len(O.strip_inline(r[c])) for r in rows))) for c in range(ncols)]
        ws = [max(0.7, l / sum(lens) * total) for l in lens]
        ws = [w / sum(ws) * total for w in ws]
    size = max(8.0, theme["size"] - 1)
    for ri, row in enumerate(rows):
        is_head = header and ri == 0
        for ci, text in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.width = Inches(ws[ci])
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.space_before = Pt(1)
            add_runs(par, text, bold=is_head, size=size, color="FFFFFF" if is_head else None)
            if is_head:
                _shade(cell, theme["th_fill"])
            elif _NUMBERISH.match(text.strip() or "x"):
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if not is_head and ri % 2 == 0 and ri > 0 and theme.get("zebra"):
                _shade(cell, theme["zebra"])
        if is_head:
            _repeat_header(tbl.rows[ri])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)       # breathing room after a table


def _fix_zoom(doc) -> None:
    """python-docx's default template has <w:zoom/> without the percent attribute the schema requires."""
    from docx.oxml.ns import qn
    try:
        z = doc.settings.element.find(qn("w:zoom"))
        if z is not None and z.get(qn("w:percent")) is None:
            z.set(qn("w:percent"), "100")
    except Exception:                                            # noqa: BLE001
        return


def _page_number_footer(doc) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    for s in doc.sections:
        par = s.footer.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for kind, txt in (("begin", None), (None, "PAGE"), ("separate", None), ("end", None)):   # one run per field part, as Word writes it
            run = par.add_run()
            run.font.size = Pt(9)
            if kind:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), kind)
            else:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = f" {txt} "
            run._r.append(e)
            if kind == "separate":
                t = OxmlElement("w:t")
                t.text = "1"
                run._r.append(t)


def render_blocks(doc, blocks: List[Dict[str, Any]], theme: Dict[str, Any], images: Dict[int, Path]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Inches, Pt
    align = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
             "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
    for bi, b in enumerate(blocks):
        t = b["type"]
        if t == "title":
            p = doc.add_paragraph(style="Title")
            add_runs(p, b["text"])
            if theme["rule_h1"]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif t == "subtitle":
            p = doc.add_paragraph(style="Subtitle")
            add_runs(p, b["text"])
            if theme["rule_h1"]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif t == "heading":
            p = doc.add_paragraph(style=f"Heading {b['level']}")
            add_runs(p, b["text"])
        elif t == "para":
            p = doc.add_paragraph()
            add_runs(p, b["text"], bold=b.get("bold", False), italic=b.get("italic", False))
            if b.get("align") in align:
                p.alignment = align[b["align"]]
            elif theme["rule_h1"] and bi > 0 and blocks[bi - 1]["type"] in ("title", "subtitle"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER          # resume: the contact line under the name
        elif t in ("bullets", "numbered"):
            base = "List Bullet" if t == "bullets" else "List Number"
            num_id = None
            for it in b["items"]:
                lvl = it.get("level", 0)
                name = base if lvl == 0 else f"{base} {lvl + 1}"
                try:
                    p = doc.add_paragraph(style=name)
                except KeyError:
                    p = doc.add_paragraph(style=base)
                add_runs(p, it["text"])
                if t == "numbered" and lvl == 0:
                    if num_id is None:
                        num_id = _new_restarted_num(doc, p) or ""
                    if num_id:
                        _use_num(p, num_id)
        elif t == "table":
            add_table(doc, b["rows"], b["header"], b.get("widths"), theme)
        elif t == "image":
            img = images[bi]
            w_px, h_px = O.image_size(img)
            avail = _content_width(doc)
            width = min(float(b["width"]) if b.get("width") else min(avail, 6.0), avail)
            if width * h_px / max(1, w_px) > 8.0:             # never taller than a page
                width = 8.0 * w_px / max(1, h_px)
            doc.add_picture(str(img), width=Inches(width))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if b.get("caption"):
                c = doc.add_paragraph()
                c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(c, b["caption"], italic=True, size=max(8, theme["size"] - 2), color="595959")
        elif t == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _prepare_images(blocks: List[Dict[str, Any]], session_id: Optional[str]) -> Dict[int, Path]:
    out: Dict[int, Path] = {}
    for i, b in enumerate(blocks):
        if b["type"] == "image":
            out[i] = O.image_file(b.get("path"), session_id, f"block {i + 1}")
    return out


# ------------------------------------------------------------------------------------------- docx_create
def _create(args: dict, ctx: ToolContext) -> str:
    try:
        docx = O.need("docx", "python-docx")
        target = O.target_path(args.get("path"), (".docx",), must_exist=False, what="a Word file")
        blocks = O.blocks_from(args)
        style = str(args.get("style") or "report").strip().lower()
        style = _ALIASES.get(style, style)
        if style not in THEMES:
            raise O.OfficeError(f"unknown style '{style}'. Use report or resume.")
        paper = str(args.get("paper") or "letter").strip().lower()
        paper = paper if paper in PAPERS else "letter"
        theme = THEMES[style]
        images = _prepare_images(blocks, ctx.session_id)
        doc = docx.Document()
        apply_theme(doc, theme, paper)
        render_blocks(doc, blocks, theme, images)
        if theme["page_numbers"]:
            _page_number_footer(doc)
        _fix_zoom(doc)
        doc.core_properties.title = O.doc_title(blocks)
        doc.core_properties.author = "Lucy"
        replaced = O.save_atomic(target, lambda tmp: doc.save(tmp))
        return O.finish(target, f"{len(blocks)} blocks, style {style}" + (", replaced the old file" if replaced else ""), [])
    except Exception as e:                                       # noqa: BLE001
        return O.fail(e)


# ------------------------------------------------------------------------------------------- docx_edit
def _iter_paragraphs(doc):
    """Every paragraph: body, table cells (nested too), headers and footers."""
    def in_table(tbl):
        for row in tbl.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                for p in cell.paragraphs:
                    yield p
                for t in cell.tables:
                    yield from in_table(t)
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        yield from in_table(t)
    for s in doc.sections:
        for part in (s.header, s.footer):
            for p in part.paragraphs:
                yield p
            for t in part.tables:
                yield from in_table(t)


def _replace_in_paragraph(par, find: str, repl: str, flags: int) -> int:
    runs = [r for r in par.runs]
    full = "".join(r.text for r in runs)
    matches = list(re.finditer(re.escape(find), full, flags))
    for m in reversed(matches):
        s, e = m.span()
        pos, first = 0, True
        for r in runs:
            txt = r.text
            rs, re_ = pos, pos + len(txt)
            pos = re_
            if not txt or re_ <= s or rs >= e:
                continue
            prefix = txt[:max(0, s - rs)]
            suffix = txt[e - rs:] if e < re_ else ""
            r.text = prefix + (repl if first else "") + suffix
            first = False
    return len(matches)


def _find_para(doc, match: str, where: str):
    """First body-level paragraph containing `match` (exact case first, then ignoring case)."""
    m = str(match or "").strip()
    if not m:
        raise O.OfficeError(f"{where}: 'match' is empty. Give a few words from the paragraph.")
    paras = doc.paragraphs
    for p in paras:
        if m in p.text:
            return p
    low = m.lower()
    for p in paras:
        if low in p.text.lower():
            return p
    in_tables = any(m.lower() in c.text.lower() for t in doc.tables for r in t.rows for c in r.cells)
    hint = " That text is inside a table: use set_cell / add_row." if in_tables else " Use office_read to copy the exact words."
    raise O.OfficeError(f"{where}: no paragraph contains '{m[:60]}'.{hint}")


def _insert_blocks(doc, anchor, blocks, theme, images, after: bool) -> None:
    body = doc.element.body
    old = list(body)
    old_ids = {id(c) for c in old}
    render_blocks(doc, blocks, theme, images)
    new = [c for c in body if id(c) not in old_ids and not c.tag.endswith("}sectPr")]
    ref = anchor._p
    if after:
        for el in new:
            ref.addnext(el)
            ref = el
    else:
        for el in new:
            ref.addprevious(el)


def _theme_of(doc) -> Dict[str, Any]:
    """Pick the house theme that matches the file (resume margins -> resume), else report."""
    try:
        return THEMES["resume"] if doc.sections[0].left_margin.inches < 0.8 else THEMES["report"]
    except Exception:                                            # noqa: BLE001
        return THEMES["report"]


def _apply_ops(doc, ops: list, session_id: Optional[str]) -> List[str]:
    theme = _theme_of(doc)
    done: List[str] = []
    for n, op in enumerate(ops, 1):
        where = f"op {n}"
        if not isinstance(op, dict):
            raise O.OfficeError(f"{where}: each op must be an object like {{\"op\": \"replace\", \"find\": \"...\", \"replace\": \"...\"}}.")
        kind = str(op.get("op") or op.get("action") or "").strip().lower().replace("-", "_")
        kind = {"replace_text": "replace", "find_replace": "replace", "add": "append", "add_paragraph": "append",
                "insert": "insert_after", "remove": "delete", "set": "set_cell", "update_cell": "set_cell",
                "add_table_row": "add_row", "append_row": "add_row"}.get(kind, kind)
        if kind == "replace":
            find = op.get("find", op.get("old", op.get("text")))
            repl = op.get("replace", op.get("new", op.get("with", "")))
            if find in (None, ""):
                raise O.OfficeError(f"{where}: replace needs 'find' (the text to look for).")
            find, repl = str(find), "" if repl is None else str(repl)
            total = 0
            for flags in (0, re.IGNORECASE):                      # exact case first, then forgiving
                paras = list(_iter_paragraphs(doc))
                for p in paras:
                    c = _replace_in_paragraph(p, find, repl, flags)
                    total += c
                    if c and op.get("all") is False:
                        break
                if total:
                    break
            if not total:
                raise O.OfficeError(f"{where}: '{find[:60]}' was not found. Use office_read to copy the exact text (it may be split by formatting: try a shorter piece).")
            done.append(f"replaced {total}x")
        elif kind in ("append", "insert_after", "insert_before"):
            blocks = O.blocks_from(op)
            images = _prepare_images(blocks, session_id)
            if kind == "append":
                render_blocks(doc, blocks, theme, images)
            else:
                anchor = _find_para(doc, op.get("match", op.get("after", op.get("before", ""))), where)
                _insert_blocks(doc, anchor, blocks, theme, images, after=(kind == "insert_after"))
            done.append(f"{kind} {len(blocks)} block(s)")
        elif kind == "delete":
            removed = 0
            while True:
                try:
                    p = _find_para(doc, op.get("match", op.get("find", op.get("text", ""))), where)
                except O.OfficeError:
                    if removed:
                        break
                    raise
                p._element.getparent().remove(p._element)
                removed += 1
                if not op.get("all"):
                    break
            done.append(f"deleted {removed} paragraph(s)")
        elif kind in ("set_cell", "add_row"):
            tables = doc.tables
            try:
                ti = int(op.get("table", 1))
            except (TypeError, ValueError):
                ti = 1
            if not 1 <= ti <= len(tables):
                raise O.OfficeError(f"{where}: table {ti} does not exist; the document has {len(tables)} table(s) (numbered from 1).")
            tbl = tables[ti - 1]
            if kind == "set_cell":
                try:
                    r, c = int(op.get("row")), int(op.get("col"))
                except (TypeError, ValueError):
                    raise O.OfficeError(f"{where}: set_cell needs numeric 'row' and 'col' (from 1).")
                if not (1 <= r <= len(tbl.rows) and 1 <= c <= len(tbl.columns)):
                    raise O.OfficeError(f"{where}: table {ti} has {len(tbl.rows)} row(s) and {len(tbl.columns)} column(s).")
                cell = tbl.cell(r - 1, c - 1)
                par = cell.paragraphs[0]
                keep = par.runs[0] if par.runs else None
                size = keep.font.size.pt if keep is not None and keep.font.size else None
                bold = bool(keep.bold) if keep is not None else False
                for extra in cell.paragraphs[1:]:
                    extra._element.getparent().remove(extra._element)
                for run in list(par.runs):
                    run._element.getparent().remove(run._element)
                add_runs(par, str(op.get("text", "")), bold=bold, size=size)
                done.append(f"table {ti} cell ({r},{c}) set")
            else:
                cells = op.get("cells", op.get("row"))
                if isinstance(cells, str):
                    cells = [x.strip() for x in cells.strip().strip("|").split("|")]
                if not isinstance(cells, list) or not cells:
                    raise O.OfficeError(f"{where}: add_row needs 'cells': a list with one text per column.")
                ncols = len(tbl.columns)
                cells = [str(x) for x in cells][:ncols] + [""] * max(0, ncols - len(cells))
                prev = tbl.rows[-1]
                row = tbl.add_row()
                for ci, text in enumerate(cells):
                    old = prev.cells[ci]
                    pr = old.paragraphs[0]
                    keep = pr.runs[0] if pr.runs else None
                    size = keep.font.size.pt if keep is not None and keep.font.size else None
                    par = row.cells[ci].paragraphs[0]
                    par.alignment = pr.alignment
                    par.paragraph_format.space_after = pr.paragraph_format.space_after
                    par.paragraph_format.space_before = pr.paragraph_format.space_before
                    add_runs(par, text, size=size)
                done.append(f"table {ti} row added")
        else:
            raise O.OfficeError(f"{where}: unknown op '{kind}'. Use replace, append, insert_after, insert_before, delete, set_cell or add_row.")
    return done


def _edit(args: dict, ctx: ToolContext) -> str:
    try:
        docx = O.need("docx", "python-docx")
        target = O.target_path(args.get("path"), (".docx",), must_exist=True, what="a Word file")
        ops = O._as_list(args.get("ops"), "ops") if args.get("ops") not in (None, "", []) else []
        if not ops:
            raise O.OfficeError("'ops' is empty. Example: [{\"op\": \"replace\", \"find\": \"2024\", \"replace\": \"2025\"}].")
        try:
            doc = docx.Document(str(target))
        except Exception as e:                                   # noqa: BLE001
            raise O.OfficeError(f"{target.name} could not be opened as a Word file (damaged, or an old .doc renamed?).") from e
        done = _apply_ops(doc, ops, ctx.session_id)
        O.save_atomic(target, lambda tmp: doc.save(tmp))
        return O.finish(target, "; ".join(done), [])
    except Exception as e:                                       # noqa: BLE001
        return O.fail(e)


TOOLS = [
    {"schema": SCHEMA_CREATE, "handler": _create, "tier": "write", "label": "write"},
    {"schema": SCHEMA_EDIT, "handler": _edit, "tier": "write", "label": "edit"},
]
