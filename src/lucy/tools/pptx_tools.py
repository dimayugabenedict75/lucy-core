"""PowerPoint decks: pptx_create and pptx_edit (python-pptx, no PowerPoint needed).

Input is structured `slides` (title / subtitle / bullets / text / left+right / image / chart / table / notes) or, easiest for a
small model, `markdown` (slides separated by a line with ---). The house "deck" theme is 16:9 with a dark title slide, a thin
accent line under each title and slide numbers. Real placeholders are used (outline view and screen readers work) and the
layouts are fixed to 16:9, so slides the user adds later in PowerPoint look right too.

pptx_edit also works on decks made elsewhere: those keep their own theme (no colors / fonts are forced on them).
"""
from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from . import _office as O
from .registry import ToolContext

MARKER = "lucy-deck"                      # core_properties.comments: this deck was made by pptx_create (our theme applies)
W_IN, H_IN = 13.333, 7.5
DECK = dict(font="Calibri", dark="1F3864", accent="2E75B6", text="262626", muted="7F7F7F", light="F3F6FA", on_dark="FFFFFF",
            sub_on_dark="D6DCE5", palette=["2E75B6", "1F3864", "ED7D31", "70AD47", "FFC000", "7F7F7F", "5B9BD5", "A5A5A5"])
KEEP_LAYOUTS = ("Title Slide", "Title and Content", "Section Header", "Two Content", "Title Only", "Blank")
LAYOUT_NAMES = ("title", "section", "bullets", "two_column", "image", "chart", "table")

SCHEMA_CREATE = {
    "type": "function",
    "function": {
        "name": "pptx_create",
        "description": ("Create a PowerPoint deck (.pptx) from 'markdown' (easiest: slides separated by a --- line) or 'slides'. "
                        "The file is sent to the chat."),
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Where to save it, e.g. C:/Users/you/Documents/pitch.pptx"},
                "markdown": {"type": "string",
                             "description": ("One block per slide, blocks separated by a line '---'. In each: '# Title', then - bullets, "
                                             "paragraphs, | tables |, ![caption](image path), and 'Notes: speaker notes'. The first slide becomes the title slide.")},
                "slides": {
                    "type": "array",
                    "description": "Instead of markdown: the slides in order.",
                    "items": {
                        "type": "object",
                        "properties": {
                            "layout": {"type": "string", "enum": list(LAYOUT_NAMES), "description": "Optional: it is guessed from the other fields."},
                            "title": {"type": "string"},
                            "subtitle": {"type": "string", "description": "title / section slides."},
                            "bullets": {"type": "array", "items": {"type": "string"}},
                            "text": {"type": "string", "description": "A plain paragraph instead of bullets."},
                            "left": {"type": "array", "items": {"type": "string"}, "description": "two_column: left bullets."},
                            "right": {"type": "array", "items": {"type": "string"}, "description": "two_column: right bullets."},
                            "image": {"type": "string", "description": "Path of a picture."},
                            "caption": {"type": "string"},
                            "chart": {"type": "object", "description": "{type: bar|line|pie|column|area, categories: [...], series: [{name, values: [...]}]}"},
                            "table": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}, "description": "Rows of cells, first row = header."},
                            "notes": {"type": "string", "description": "Speaker notes."},
                        },
                    },
                },
            },
            "required": ["path"],
        },
    },
}

SCHEMA_EDIT = {
    "type": "function",
    "function": {
        "name": "pptx_edit",
        "description": ("Change an existing PowerPoint deck with a list of 'ops' (all or none). Read it first with office_read to see "
                        "the slide numbers (they start at 1)."),
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "The .pptx file (must exist)."},
                "ops": {
                    "type": "array",
                    "description": ("add_slide {slide:{title, bullets|image|chart|table|..}, position?}; delete_slide {slide}; move_slide {slide, to}; "
                                    "replace_text {find, replace, slide?}; set_text {slide, title?, bullets?, notes?}."),
                    "items": {
                        "type": "object",
                        "properties": {
                            "op": {"type": "string", "enum": ["add_slide", "delete_slide", "move_slide", "replace_text", "set_text"]},
                            "slide": {"type": ["integer", "object"], "description": "Slide number; for add_slide the new slide's fields."},
                            "position": {"type": "integer", "description": "add_slide: where it goes (default: the end)."},
                            "to": {"type": "integer"}, "find": {"type": "string"}, "replace": {"type": "string"},
                            "title": {"type": "string"}, "bullets": {"type": "array", "items": {"type": "string"}},
                            "notes": {"type": "string"},
                        },
                        "required": ["op"],
                    },
                },
            },
            "required": ["path", "ops"],
        },
    },
}


def _pp():
    pptx = O.need("pptx", "python-pptx")
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Emu, Inches, Pt
    return pptx, RGBColor, MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN, Emu, Inches, Pt


def _rgb(h: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(h.upper())


# ------------------------------------------------------------------------------------------- slide specs
def _s(v: Any) -> str:
    return "" if v is None else str(v)


def norm_slide(spec: Any, i: int) -> Dict[str, Any]:
    where = f"slide {i}"
    if isinstance(spec, str):
        spec = {"title": spec}
    if not isinstance(spec, dict):
        raise O.OfficeError(f"{where}: each slide must be an object like {{\"title\": \"...\", \"bullets\": [\"...\"]}}.")
    s: Dict[str, Any] = {"title": _s(spec.get("title", spec.get("heading", ""))).strip(), "subtitle": _s(spec.get("subtitle")).strip(),
                         "notes": _s(spec.get("notes", spec.get("speaker_notes"))).strip(), "caption": _s(spec.get("caption")).strip(),
                         "text": _s(spec.get("text", spec.get("body", spec.get("content") if isinstance(spec.get("content"), str) and not spec.get("bullets") else ""))).strip()}
    for key in ("bullets", "left", "right"):
        v = spec.get(key, spec.get("points") if key == "bullets" else None)
        s[key] = O._items(v) if v not in (None, "", []) else []
    img = spec.get("image", spec.get("img", spec.get("picture")))
    s["image"] = _s(img.get("path") if isinstance(img, dict) else img).strip()
    tb = spec.get("table")
    if isinstance(tb, dict):
        tb = tb.get("rows", tb.get("data"))
    try:
        s["table"] = O._rows(tb) if tb not in (None, "", []) else []
    except O.OfficeError as e:
        raise O.OfficeError(f"{where}: {e}") from e
    ch = spec.get("chart")
    if isinstance(ch, str):
        try:
            import json
            ch = json.loads(ch)
        except ValueError:
            raise O.OfficeError(f"{where}: 'chart' must be an object: {{\"type\": \"bar\", \"categories\": [...], \"series\": [{{\"name\": \"..\", \"values\": [...]}}]}}.")
    s["chart"] = _norm_chart(ch, where) if ch else None
    lay = str(spec.get("layout") or "").strip().lower().replace(" ", "_").replace("-", "_")
    lay = {"title_slide": "title", "cover": "title", "section_header": "section", "divider": "section", "content": "bullets", "list": "bullets",
           "title_and_content": "bullets", "two_content": "two_column", "two_columns": "two_column", "picture": "image", "photo": "image",
           "graph": "chart"}.get(lay, lay)
    if lay not in LAYOUT_NAMES:
        lay = ""
    if not lay:
        lay = ("image" if s["image"] else "chart" if s["chart"] else "table" if s["table"] else "two_column" if (s["left"] or s["right"])
               else "bullets" if (s["bullets"] or s["text"]) else "title" if i == 1 else "section")
    if lay == "bullets" and not (s["bullets"] or s["text"]) and not s["title"]:
        raise O.OfficeError(f"{where}: empty slide. Give a title and 'bullets' (or an image / chart / table).")
    if not s["title"] and lay in ("title", "section", "image", "chart", "table"):
        raise O.OfficeError(f"{where}: a '{lay}' slide needs a 'title'.")
    s["layout"] = lay
    return s


def _norm_chart(ch: Any, where: str) -> Dict[str, Any]:
    if not isinstance(ch, dict):
        raise O.OfficeError(f"{where}: 'chart' must be an object with type, categories and series.")
    kind = str(ch.get("type", "bar")).strip().lower().replace(" ", "_")
    kind = {"column": "column", "bar": "column", "vertical_bar": "column", "bars": "column", "barh": "bar", "horizontal_bar": "bar",
            "line": "line", "lines": "line", "pie": "pie", "doughnut": "doughnut", "donut": "doughnut", "area": "area",
            "stacked": "stacked", "stacked_bar": "stacked", "stacked_column": "stacked"}.get(kind, kind)
    if kind not in ("column", "bar", "line", "pie", "doughnut", "area", "stacked"):
        raise O.OfficeError(f"{where}: chart type '{kind}' is not supported. Use bar, line, pie, area or stacked.")
    cats = ch.get("categories", ch.get("labels", ch.get("x")))
    series = ch.get("series")
    if series is None and ch.get("values") is not None:
        series = [{"name": ch.get("name", ""), "values": ch.get("values")}]
    if isinstance(series, dict):                                   # {"Sales": [1,2,3]} -> list form
        series = [{"name": k, "values": v} for k, v in series.items()]
    if not isinstance(cats, list) or not cats or not isinstance(series, list) or not series:
        raise O.OfficeError(f"{where}: a chart needs 'categories' (list of labels) and 'series' ([{{\"name\": \"..\", \"values\": [numbers]}}]).")
    out = []
    for k, sr in enumerate(series, 1):
        if not isinstance(sr, dict):
            raise O.OfficeError(f"{where}: series {k} must be {{\"name\": \"..\", \"values\": [numbers]}}.")
        vals = sr.get("values", sr.get("data"))
        try:
            nums = [float(str(v).replace(",", "").replace("%", "")) for v in vals]
        except (TypeError, ValueError):
            raise O.OfficeError(f"{where}: series {k} 'values' must be a list of numbers.")
        if len(nums) != len(cats):
            raise O.OfficeError(f"{where}: series {k} has {len(nums)} values but there are {len(cats)} categories. They must match.")
        out.append({"name": _s(sr.get("name", f"Series {k}")), "values": nums})
    return {"type": kind, "categories": [_s(c) for c in cats], "series": out, "number_format": _s(ch.get("number_format"))}


_BREAK = re.compile(r"^\s*-{3,}\s*$", re.M)
_NOTES = re.compile(r"^\s*(?:speaker\s+)?notes?\s*:\s*(.*)$", re.I)


def parse_deck_markdown(md: str) -> List[Dict[str, Any]]:
    """Markdown deck -> slide specs. Slides are separated by a bare '---' line; the first heading of a slide is its title."""
    slides: List[Dict[str, Any]] = []
    for chunk in _BREAK.split(str(md or "").replace("\r\n", "\n")):
        if not chunk.strip():
            continue
        notes: List[str] = []
        keep: List[str] = []
        in_notes = False
        for ln in chunk.split("\n"):
            m = _NOTES.match(ln)
            if m:
                in_notes = True
                notes.append(m.group(1))
            elif in_notes and ln.strip():
                notes.append(ln.strip())
            else:
                in_notes = in_notes and bool(ln.strip())
                keep.append(ln)
        blocks = O.parse_markdown("\n".join(keep))
        s: Dict[str, Any] = {"bullets": [], "notes": " ".join(n for n in notes if n).strip()}
        paras: List[str] = []
        for b in blocks:
            t = b["type"]
            if t in ("title", "heading") and not s.get("title"):
                s["title"] = b["text"]
            elif t in ("title", "heading", "para"):
                paras.append(b["text"])
            elif t in ("bullets", "numbered"):
                s["bullets"] += [{"text": (f"{n}. " if t == "numbered" and it["level"] == 0 else "") + it["text"], "level": it["level"]}
                                 for n, it in enumerate(b["items"], 1)]
            elif t == "table":
                s["table"] = b["rows"]
            elif t == "image":
                s["image"], s["caption"] = b["path"], b.get("caption", "")
        if paras:
            if not s["bullets"] and len(slides) == 0 and "title" in s:
                s["subtitle"] = " ".join(paras)
            elif not s["bullets"] and "image" not in s and "table" not in s:
                s["text"] = "\n".join(paras)
            else:
                s["bullets"] = [{"text": p, "level": 0} for p in paras] + s["bullets"]
        if not s["bullets"]:
            s.pop("bullets")
        slides.append(s)
    if not slides:
        raise O.OfficeError("no slides found in 'markdown'. Write '# Title' and bullets for each slide, separated by a --- line.")
    return slides


def slides_from(args: dict) -> List[Dict[str, Any]]:
    sl, md = args.get("slides"), args.get("markdown")
    if sl not in (None, "", []):
        raw = O._as_list(sl, "slides")
    elif isinstance(md, str) and md.strip():
        raw = parse_deck_markdown(md)
    else:
        raise O.OfficeError("nothing to write: give 'markdown' (slides separated by ---) or 'slides' (list).")
    return [norm_slide(s, i) for i, s in enumerate(raw, 1)]


# ------------------------------------------------------------------------------------------- geometry + text
def _scale(prs):
    from pptx.util import Inches
    return prs.slide_width / Inches(W_IN), prs.slide_height / Inches(H_IN)


def _box(prs, x: float, y: float, w: float, h: float):
    from pptx.util import Inches
    sx, sy = _scale(prs)
    return int(Inches(x) * sx), int(Inches(y) * sy), int(Inches(w) * sx), int(Inches(h) * sy)


def _fit(paras: List[Tuple[str, int]], width_in: float, height_in: float, start: int = 24, lo: int = 14) -> Tuple[int, bool]:
    """Biggest font (start..lo, step 2) at which the paragraphs probably fit the box. (size, fits?)"""
    for size in range(start, lo - 1, -2):
        total = 0.0
        for text, level in paras:
            sz = max(10, size - 2 * level)
            w = max(1.0, width_in - 0.5 - 0.45 * level)
            cpl = max(8, int(w * 72 / (sz * 0.52)))
            lines = max(1, math.ceil(len(O.strip_inline(text)) / cpl))
            total += lines * sz * 1.2 + sz * 0.45
        if total / 72.0 <= height_in * 0.95:
            return size, True
    return lo, False


def _write_runs(par, text: str, size: Optional[float], color: Optional[str], font: Optional[str], bold: Optional[bool] = None) -> None:
    from pptx.util import Pt
    for seg, b, i, code in O.parse_inline(text):
        if seg == "":
            continue
        r = par.add_run()
        r.text = seg
        if size:
            r.font.size = Pt(size)
        if font:
            r.font.name = "Consolas" if code else font
        if color:
            r.font.color.rgb = _rgb(color)
        if b or bold:
            r.font.bold = True
        if i:
            r.font.italic = True


def _fill(tf, paras: List[Tuple[str, int]], size: float, st: Optional[dict], *, color: Optional[str] = None, bold: Optional[bool] = None,
          bullets: bool = True, align=None) -> None:
    from pptx.enum.text import MSO_AUTO_SIZE
    tf.word_wrap = True
    tf.clear()
    for k, (text, level) in enumerate(paras):
        par = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        par.level = max(0, min(4, level))
        if align is not None:
            par.alignment = align
        sz = max(10, size - 2 * level) if st else None
        _write_runs(par, text, sz, color or (st["text"] if st else None), st["font"] if st else None, bold)
        if not bullets:
            ppr = par._p.get_or_add_pPr()
            ppr.set("marL", "0")
            ppr.set("indent", "0")
            from pptx.oxml.ns import qn
            for old in ppr.findall(qn("a:buNone")):
                ppr.remove(old)
            from lxml import etree
            etree.SubElement(ppr, qn("a:buNone"))
        if st:
            par.space_after = _pt(6 if level == 0 else 3)
    if st:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _pt(v: float):
    from pptx.util import Pt
    return Pt(v)


def _rect(slide, x, y, w, h, fill: str):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = _rgb(fill)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _slide_number(prs, slide, st: dict) -> None:
    """A real slide-number field (it renumbers itself when slides move)."""
    from lxml import etree
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    x, y, w, h = _box(prs, 11.9, 6.95, 0.9, 0.35)
    tb = slide.shapes.add_textbox(x, y, w, h)
    tb.name = "Slide Number"
    par = tb.text_frame.paragraphs[0]
    par.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(par._p, qn("a:fld"))
    fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
    fld.set("type", "slidenum")
    rpr = etree.SubElement(fld, qn("a:rPr"))
    rpr.set("lang", "en-US")
    rpr.set("sz", "1100")
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr")).set("val", st["muted"])
    etree.SubElement(fld, qn("a:t")).text = "\u2039#\u203a"


# ------------------------------------------------------------------------------------------- deck + layouts
def _layout(prs, name: str):
    for lay in prs.slide_layouts:
        if lay.name.strip().lower() == name.lower():
            return lay
    idx = {"Title Slide": 0, "Title and Content": 1, "Section Header": 2, "Two Content": 3, "Title Only": 5, "Blank": 6}.get(name, 1)
    return prs.slide_layouts[min(idx, len(prs.slide_layouts) - 1)]


def _place(shape, geo) -> None:
    shape.left, shape.top, shape.width, shape.height = geo


def new_deck():
    """16:9 deck with the unused 4:3 layouts removed and the used ones re-positioned (so 'New slide' in PowerPoint looks right)."""
    pptx, *_ = _pp()
    from pptx.util import Inches
    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN)
    for lay in list(prs.slide_layouts):
        if lay.name not in KEEP_LAYOUTS:
            try:
                prs.slide_layouts.remove(lay)
            except Exception:                                      # noqa: BLE001
                pass
    t_geo, b_geo = _box(prs, 0.7, 0.4, 11.93, 0.95), _box(prs, 0.7, 1.75, 11.93, 5.0)
    for ph in prs.slide_master.placeholders:
        k = ph.placeholder_format.type
        if "TITLE" in str(k):
            _place(ph, t_geo)
        elif "BODY" in str(k):
            _place(ph, b_geo)
    for lay in prs.slide_layouts:
        for ph in lay.placeholders:
            idx, kind = ph.placeholder_format.idx, str(ph.placeholder_format.type)
            if lay.name == "Title Slide":
                _place(ph, _box(prs, 0.9, 2.75, 11.5, 1.5) if idx == 0 else _box(prs, 0.9, 4.2, 11.5, 1.2)) if idx in (0, 1) else None
            elif lay.name == "Section Header":
                _place(ph, _box(prs, 0.9, 2.9, 11.5, 1.4) if idx == 0 else _box(prs, 0.9, 4.3, 11.5, 1.0)) if idx in (0, 1) else None
            elif lay.name == "Two Content":
                if idx == 0:
                    _place(ph, t_geo)
                elif idx == 1:
                    _place(ph, _box(prs, 0.7, 1.75, 5.8, 5.0))
                elif idx == 2:
                    _place(ph, _box(prs, 6.83, 1.75, 5.8, 5.0))
            elif idx == 0:
                _place(ph, t_geo)
            elif idx == 1 and "BODY" in kind or idx == 1 and "OBJECT" in kind:
                _place(ph, b_geo)
    prs.core_properties.comments = MARKER
    prs.core_properties.author = "Lucy"
    return prs


def _is_ours(prs) -> bool:
    return (prs.core_properties.comments or "").strip() == MARKER


def _drop(shape) -> None:
    el = shape._element
    el.getparent().remove(el)


def _placeholders(slide) -> Dict[int, Any]:
    return {p.placeholder_format.idx: p for p in slide.placeholders}


def _title_text(slide, prs, text: str, st: Optional[dict], on_dark: bool = False, size: int = 30) -> None:
    from pptx.enum.text import MSO_ANCHOR
    t = slide.shapes.title
    if t is None:
        return
    if st:
        n = len(text)
        size = size if n <= 55 else max(20, size - 6) if n <= 85 else 18
    from pptx.enum.text import PP_ALIGN
    _fill(t.text_frame, [(text, 0)], size, st, color=(st["on_dark"] if on_dark else st["dark"]) if st else None, bold=True if st else None,
          bullets=False, align=PP_ALIGN.LEFT if st else None)
    if st:
        t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE if not on_dark else MSO_ANCHOR.TOP


# ------------------------------------------------------------------------------------------- slide builders
def _content_box(prs, slide, st: Optional[dict]):
    """Where a picture / chart / table goes: below our title (our deck) or inside the layout's body box (foreign deck)."""
    if st:
        return _box(prs, 0.7, 1.75, 11.93, 5.0)
    body = [p for p in slide.placeholders if p.placeholder_format.idx == 1]
    if body:
        g = (body[0].left, body[0].top, body[0].width, body[0].height)
        _drop(body[0])
        return g
    return _box(prs, 0.7, 1.75, 11.93, 5.0)


def _add_picture(slide, path: Path, geo, caption: str, st: Optional[dict], prs) -> None:
    from pptx.enum.text import PP_ALIGN
    x, y, w, h = geo
    cap_h = int(h * 0.09) if caption else 0
    pw, ph = O.image_size(path)
    scale = min(w / pw, (h - cap_h) / ph)
    dw, dh = int(pw * scale), int(ph * scale)
    slide.shapes.add_picture(str(path), x + (w - dw) // 2, y + ((h - cap_h) - dh) // 2, dw, dh)
    if caption:
        tb = slide.shapes.add_textbox(x, y + h - cap_h, w, cap_h)
        _fill(tb.text_frame, [(caption, 0)], 14, st or {"text": "595959", "font": "Calibri"}, color="595959", bullets=False, align=PP_ALIGN.CENTER)


def _add_chart(slide, ch: Dict[str, Any], geo, st: Optional[dict]) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Pt
    kinds = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
             "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT, "area": XL_CHART_TYPE.AREA, "stacked": XL_CHART_TYPE.COLUMN_STACKED}
    data = CategoryChartData()
    data.categories = ch["categories"]
    for s in ch["series"]:
        data.add_series(s["name"], s["values"])
    x, y, w, h = geo
    gf = slide.shapes.add_chart(kinds[ch["type"]], x, y, w, h, data)
    chart = gf.chart
    chart.font.size = Pt(13)
    if st:
        chart.font.name = st["font"]
    chart.has_title = False
    pal = (st or DECK)["palette"] if st else DECK["palette"]
    pie = ch["type"] in ("pie", "doughnut")
    chart.has_legend = pie or len(ch["series"]) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    plot = chart.plots[0]
    if pie:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_percentage, dl.show_value = True, False
        dl.number_format, dl.number_format_is_linked = "0%", False
        dl.font.size, dl.font.bold = Pt(15), True
        dl.font.color.rgb = _rgb("FFFFFF")
        for k in range(len(ch["categories"])):
            pt = plot.series[0].points[k]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _rgb(pal[k % len(pal)])
    else:
        for k, s in enumerate(plot.series):
            col = _rgb(pal[k % len(pal)])
            if ch["type"] == "line":
                s.format.line.color.rgb = col
                s.format.line.width = Pt(2.5)
                s.smooth = False
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = col
        try:
            chart.value_axis.major_gridlines.format.line.color.rgb = _rgb("D9D9D9")
            chart.value_axis.format.line.fill.background()
            if ch.get("number_format"):
                chart.value_axis.tick_labels.number_format = ch["number_format"]
                chart.value_axis.tick_labels.number_format_is_linked = False
        except Exception:                                          # noqa: BLE001
            pass


def _add_table(slide, rows: List[List[str]], geo, st: Optional[dict]) -> None:
    from pptx.enum.text import PP_ALIGN
    x, y, w, h = geo
    n, c = len(rows), len(rows[0])
    size = 18 if n <= 6 else 15 if n <= 10 else 12
    row_h = min(int(h / n), int(size * 2.2 * 12700))
    gf = slide.shapes.add_table(n, c, x, y, w, row_h * n)
    tbl = gf.table
    for ri, row in enumerate(rows):
        for ci, text in enumerate(row):
            cell = tbl.cell(ri, ci)
            head = ri == 0
            _fill(cell.text_frame, [(text, 0)], size, st or {"text": "262626", "font": "Calibri"},
                  color="FFFFFF" if head and st else None, bold=True if head else None, bullets=False,
                  align=PP_ALIGN.RIGHT if (not head and re.fullmatch(r"[\(\-+]?\s*[$\u20ac\u00a3\u20b1]?\s*[\d,]+(\.\d+)?\s*%?\)?", text.strip() or "x")) else None)
            if st:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(st["dark"] if head else (st["light"] if ri % 2 == 0 else "FFFFFF"))


def build_slide(prs, spec: Dict[str, Any], st: Optional[dict], images: Dict[str, Path], n: int) -> List[str]:
    """Add one slide at the end. Returns problems (text that cannot fit...)."""
    from pptx.enum.text import PP_ALIGN
    problems: List[str] = []
    lay = spec["layout"]
    name = {"title": "Title Slide", "section": "Section Header", "bullets": "Title and Content", "two_column": "Two Content"}.get(lay, "Title Only")
    slide = prs.slides.add_slide(_layout(prs, name))
    phs = _placeholders(slide)
    if lay in ("title", "section"):
        if st:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = _rgb(st["dark"])
            x, y, w, h = _box(prs, 0.9, 2.55 if lay == "title" else 2.7, 1.4, 0.09)
            _rect(slide, x, y, w, h, st["accent"])
        _title_text(slide, prs, spec["title"], st, on_dark=True, size=44 if lay == "title" else 40)
        sub = phs.get(1)
        if sub is not None:
            if spec["subtitle"] or spec["text"]:
                _fill(sub.text_frame, [(spec["subtitle"] or spec["text"], 0)], 22, st, color=st["sub_on_dark"] if st else None, bullets=False,
                      align=PP_ALIGN.LEFT if st else None)
            else:
                _drop(sub)
    else:
        _title_text(slide, prs, spec["title"], st)
        if st:
            x, y, w, h = _box(prs, 0.7, 1.4, 1.2, 0.06)
            _rect(slide, x, y, w, h, st["accent"])
        if lay == "bullets":
            body = phs.get(1)
            paras = [(b["text"], b["level"]) for b in spec["bullets"]] or [(t, 0) for t in spec["text"].split("\n") if t.strip()]
            size, fits = _fit(paras, 11.93, 5.0)
            if len(paras) > 12:
                fits = False
            if not fits:
                problems.append(f"Slide {n} has too much text to fit. Shorten it or split it over two slides.")
            if body is not None:
                _fill(body.text_frame, paras, size, st, bullets=bool(spec["bullets"]))
        elif lay == "two_column":
            for idx, key in ((1, "left"), (2, "right")):
                ph = phs.get(idx)
                paras = [(b["text"], b["level"]) for b in spec[key]]
                if ph is None:
                    continue
                if not paras:
                    _drop(ph)
                    continue
                size, fits = _fit(paras, 5.8, 5.0, start=22)
                if not fits or len(paras) > 10:
                    problems.append(f"Slide {n}: the {key} column has too much text. Shorten it.")
                _fill(ph.text_frame, paras, size, st)
        else:
            geo = _content_box(prs, slide, st)
            if lay == "image":
                x, y, w, h = geo
                paras = [(b["text"], b["level"]) for b in spec["bullets"]]
                if paras:                                            # picture left, short bullets right
                    pw = int(w * 0.58)
                    _add_picture(slide, images[spec["image"]], (x, y, pw, h), spec["caption"], st, prs)
                    tb = slide.shapes.add_textbox(x + pw + int(w * 0.04), y, w - pw - int(w * 0.04), h)
                    size, fits = _fit(paras, (w - pw) / 914400.0 * 0.96, h / 914400.0, start=20)
                    if not fits:
                        problems.append(f"Slide {n}: the text next to the picture is too long.")
                    _fill(tb.text_frame, paras, size, st or {"text": "262626", "font": "Calibri"})
                else:
                    _add_picture(slide, images[spec["image"]], geo, spec["caption"], st, prs)
            elif lay == "chart":
                _add_chart(slide, spec["chart"], geo, st)
            elif lay == "table":
                if len(spec["table"]) > 14:
                    problems.append(f"Slide {n}: the table has {len(spec['table'])} rows, too many for one slide. Split it.")
                _add_table(slide, spec["table"], geo, st)
    if spec["notes"]:
        slide.notes_slide.notes_text_frame.text = spec["notes"]
    if st and lay not in ("title", "section"):
        _slide_number(prs, slide, st)
    return problems


def _prepare_images(slides: List[Dict[str, Any]], session_id: Optional[str]) -> Dict[str, Path]:
    out: Dict[str, Path] = {}
    for i, s in enumerate(slides, 1):
        if s.get("image") and s["image"] not in out:
            out[s["image"]] = O.image_file(s["image"], session_id, f"slide {i}")
    return out


# ------------------------------------------------------------------------------------------- pptx_create
def _create(args: dict, ctx: ToolContext) -> str:
    try:
        O.need("pptx", "python-pptx")
        target = O.target_path(args.get("path"), (".pptx",), must_exist=False, what="a PowerPoint file")
        slides = slides_from(args)
        images = _prepare_images(slides, ctx.session_id)
        prs = new_deck()
        problems: List[str] = []
        for n, s in enumerate(slides, 1):
            problems += build_slide(prs, s, DECK, images, n)
        title = next((s["title"] for s in slides if s["title"]), "")
        prs.core_properties.title = title[:120]
        replaced = O.save_atomic(target, lambda tmp: prs.save(tmp))
        return O.finish(target, f"{len(slides)} slides" + (", replaced the old file" if replaced else ""), problems)
    except Exception as e:                                           # noqa: BLE001
        return O.fail(e)


# ------------------------------------------------------------------------------------------- pptx_edit
def _slide_no(prs, v: Any, where: str) -> int:
    try:
        n = int(v)
    except (TypeError, ValueError):
        raise O.OfficeError(f"{where}: 'slide' must be a slide number (from 1).")
    if not 1 <= n <= len(prs.slides):
        raise O.OfficeError(f"{where}: the deck has {len(prs.slides)} slide(s); there is no slide {n}.")
    return n


def _move(prs, old: int, new: int) -> None:
    lst = prs.slides._sldIdLst
    items = list(lst)
    el = items[old - 1]
    lst.remove(el)
    lst.insert(max(0, min(len(items) - 1, new - 1)), el)


def _renumber(prs) -> None:
    """slide1.xml, slide2.xml ... in deck order. Without this a slide added after a delete is named after the slide COUNT and
    collides with a slide that kept its old name: two parts with one name in the file, and PowerPoint asks to repair it."""
    prs.part.rename_slide_parts([el.rId for el in prs.slides._sldIdLst])


def _delete(prs, n: int) -> None:
    lst = prs.slides._sldIdLst
    el = list(lst)[n - 1]
    prs.part.drop_rel(el.rId)
    lst.remove(el)
    _renumber(prs)


def _text_frames(slide, with_notes: bool = False):
    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:                                   # group
                yield from walk(sh.shapes)
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                yield sh.text_frame
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame
    yield from walk(slide.shapes)
    if with_notes and slide.has_notes_slide:
        yield slide.notes_slide.notes_text_frame


def _replace_in_par(par, find: str, repl: str, flags: int) -> int:
    runs = list(par.runs)
    full = "".join(r.text for r in runs)
    ms = list(re.finditer(re.escape(find), full, flags))
    for m in reversed(ms):
        s, e = m.span()
        pos, first = 0, True
        for r in runs:
            txt = r.text
            rs, re_ = pos, pos + len(txt)
            pos = re_
            if not txt or re_ <= s or rs >= e:
                continue
            r.text = txt[:max(0, s - rs)] + (repl if first else "") + (txt[e - rs:] if e < re_ else "")
            first = False
    return len(ms)


def _apply_ops(prs, ops: list, session_id: Optional[str]) -> Tuple[List[str], List[str]]:
    done: List[str] = []
    problems: List[str] = []
    for n, op in enumerate(ops, 1):
        where = f"op {n}"
        if not isinstance(op, dict):
            raise O.OfficeError(f"{where}: each op must be an object like {{\"op\": \"delete_slide\", \"slide\": 3}}.")
        kind = str(op.get("op") or op.get("action") or "").strip().lower().replace("-", "_")
        kind = {"add": "add_slide", "insert_slide": "add_slide", "new_slide": "add_slide", "remove_slide": "delete_slide",
                "delete": "delete_slide", "move": "move_slide", "reorder": "move_slide", "replace": "replace_text", "find_replace": "replace_text",
                "set_notes": "set_text", "update_slide": "set_text", "edit_slide": "set_text"}.get(kind, kind)
        if kind == "add_slide":
            spec = op.get("slide") if isinstance(op.get("slide"), dict) else {k: v for k, v in op.items() if k not in ("op", "action", "position", "slide")}
            ns = norm_slide(spec, len(prs.slides) + 1)
            imgs = _prepare_images([ns], session_id)
            st = DECK if _is_ours(prs) else None
            problems += build_slide(prs, ns, st, imgs, len(prs.slides) + 1)
            pos = op.get("position")
            if pos not in (None, ""):
                _move(prs, len(prs.slides), max(1, min(int(pos), len(prs.slides))))
            done.append("slide added")
        elif kind == "delete_slide":
            _delete(prs, _slide_no(prs, op.get("slide"), where))
            done.append("slide deleted")
        elif kind == "move_slide":
            old = _slide_no(prs, op.get("slide"), where)
            new = _slide_no(prs, op.get("to", op.get("position")), where)
            _move(prs, old, new)
            done.append(f"slide {old} moved to {new}")
        elif kind == "replace_text":
            find, repl = op.get("find", op.get("old")), op.get("replace", op.get("new", ""))
            if find in (None, ""):
                raise O.OfficeError(f"{where}: replace_text needs 'find'.")
            find, repl = str(find), "" if repl is None else str(repl)
            sl = [prs.slides[_slide_no(prs, op["slide"], where) - 1]] if op.get("slide") not in (None, "") else list(prs.slides)
            total = 0
            for flags in (0, re.IGNORECASE):
                for slide in sl:
                    for tf in _text_frames(slide, with_notes=True):
                        for par in tf.paragraphs:
                            total += _replace_in_par(par, find, repl, flags)
                if total:
                    break
            if not total:
                raise O.OfficeError(f"{where}: '{find[:60]}' was not found. Use office_read to copy the exact text.")
            done.append(f"replaced {total}x")
        elif kind == "set_text":
            slide = prs.slides[_slide_no(prs, op.get("slide"), where) - 1]
            st = DECK if _is_ours(prs) else None
            changed = False
            if op.get("title") not in (None, ""):
                if slide.shapes.title is None:
                    raise O.OfficeError(f"{where}: slide has no title box.")
                lay = slide.slide_layout.name
                _title_text(slide, prs, str(op["title"]), st, on_dark=lay in ("Title Slide", "Section Header"),
                            size=44 if lay == "Title Slide" else 40 if lay == "Section Header" else 30)
                changed = True
            if op.get("bullets") not in (None, "", []):
                items = O._items(op["bullets"])
                body = [p for p in slide.placeholders if p.placeholder_format.idx == 1 and p.placeholder_format.type is not None and "TITLE" not in str(p.placeholder_format.type)]
                if not body:
                    raise O.OfficeError(f"{where}: slide {op.get('slide')} has no text box for bullets (it may be an image / chart slide).")
                paras = [(b["text"], b["level"]) for b in items]
                size, fits = _fit(paras, body[0].width / 914400.0, body[0].height / 914400.0)
                if not fits:
                    problems.append(f"Slide {op.get('slide')} now has too much text to fit. Shorten it or split it.")
                _fill(body[0].text_frame, paras, size, st)
                changed = True
            if op.get("notes") not in (None,):
                slide.notes_slide.notes_text_frame.text = str(op.get("notes") or "")
                changed = True
            if not changed:
                raise O.OfficeError(f"{where}: set_text needs at least one of title, bullets, notes.")
            done.append("slide text set")
        else:
            raise O.OfficeError(f"{where}: unknown op '{kind}'. Use add_slide, delete_slide, move_slide, replace_text or set_text.")
    return done, problems


def _edit(args: dict, ctx: ToolContext) -> str:
    try:
        pptx = O.need("pptx", "python-pptx")
        target = O.target_path(args.get("path"), (".pptx",), must_exist=True, what="a PowerPoint file")
        ops = O._as_list(args.get("ops"), "ops") if args.get("ops") not in (None, "", []) else []
        if not ops:
            raise O.OfficeError("'ops' is empty. Example: [{\"op\": \"delete_slide\", \"slide\": 3}].")
        try:
            prs = pptx.Presentation(str(target))
        except Exception as e:                                       # noqa: BLE001
            raise O.OfficeError(f"{target.name} could not be opened as a PowerPoint file ({type(e).__name__}). Is it a real .pptx (not .ppt)?") from e
        done, problems = _apply_ops(prs, ops, ctx.session_id)
        _renumber(prs)
        O.save_atomic(target, lambda tmp: prs.save(tmp))
        return O.finish(target, "; ".join(done), problems)
    except Exception as e:                                           # noqa: BLE001
        return O.fail(e)


TOOLS = [
    {"schema": SCHEMA_CREATE, "handler": _create, "tier": "write", "label": "write"},
    {"schema": SCHEMA_EDIT, "handler": _edit, "tier": "write", "label": "edit"},
]
