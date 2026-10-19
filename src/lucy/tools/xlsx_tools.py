"""Excel workbooks: xlsx_create and xlsx_edit (openpyxl, no Excel needed).

Input is structured: sheets with `columns` (header row) and `rows`. A cell that is a string starting with "=" is a
formula (so the sheet recalculates), numbers typed as text ("1,200") become numbers, ISO dates become dates. After saving,
if LibreOffice is installed the workbook is recalculated on a copy and any formula that shows #DIV/0!, #REF!... is reported
back to the model BEFORE the file goes to the chat.
"""
from __future__ import annotations

import datetime as _dt
import re
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from . import _office as O
from .registry import ToolContext

FONT = "Arial"
HEAD_FILL = "1F3864"
FORMATS = {
    "currency": '"$"#,##0.00', "usd": '"$"#,##0.00', "dollar": '"$"#,##0.00', "peso": '"\u20b1"#,##0.00', "php": '"\u20b1"#,##0.00',
    "euro": '"\u20ac"#,##0.00', "percent": "0.0%", "percentage": "0.0%", "pct": "0.0%", "integer": "#,##0", "int": "#,##0",
    "number": "#,##0.00", "decimal": "#,##0.00", "date": "yyyy-mm-dd", "datetime": "yyyy-mm-dd hh:mm", "text": "@", "general": "General",
}

_SHEET_PROPS = {
    "name": {"type": "string", "description": "Sheet (tab) name."},
    "columns": {"type": "array", "items": {"type": "string"}, "description": "Header row texts (shown bold on a colored band)."},
    "rows": {"type": "array", "items": {"type": "array", "items": {}},
             "description": "Data rows (lists of cells). A text starting with '=' is a formula, e.g. '=B2*C2'."},
    "formulas": {"type": "object", "description": "Extra formulas by cell, e.g. {\"D12\": \"=SUM(D2:D11)\"}."},
    "number_formats": {"type": "object", "description": "Cell formats by column or range, e.g. {\"C\": \"currency\", \"D2:D20\": \"percent\"}. Also: integer, number, date, or an Excel code."},
    "column_widths": {"type": "object", "description": "Widths by column letter, e.g. {\"A\": 24}. Default: fitted to the content."},
    "freeze": {"type": ["boolean", "string"], "description": "true = freeze the header row; or a cell like 'B2'."},
}

SCHEMA_CREATE = {
    "type": "function",
    "function": {
        "name": "xlsx_create",
        "description": ("Create an Excel workbook (.xlsx) from 'sheets'. The file is sent to the chat; use formulas ('=SUM(B2:B9)') "
                        "for anything calculated."),
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Where to save it, e.g. C:/Users/you/Documents/budget.xlsx"},
                "sheets": {"type": "array", "description": "The sheets in order.",
                           "items": {"type": "object", "properties": _SHEET_PROPS, "required": ["name"]}},
            },
            "required": ["path", "sheets"],
        },
    },
}

SCHEMA_EDIT = {
    "type": "function",
    "function": {
        "name": "xlsx_edit",
        "description": ("Change an existing Excel workbook with a list of 'ops' (all or none); formulas are kept. "
                        "Look at it first with office_read."),
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "The .xlsx file (must exist)."},
                "ops": {
                    "type": "array",
                    "description": ("set_cells {sheet, cells:{'B2': 5, 'C2': '=B2*2'}}; append_rows {sheet, rows}; add_sheet {name, columns, rows}; "
                                    "delete_sheet {sheet}; rename_sheet {sheet, name}; number_format {sheet, range, format}; "
                                    "column_width {sheet, widths}; delete_rows {sheet, row, count}; clear {sheet, range}. "
                                    "'sheet' is a name or a number from 1 (default: the first)."),
                    "items": {
                        "type": "object",
                        "properties": {
                            "op": {"type": "string", "enum": ["set_cells", "append_rows", "add_sheet", "delete_sheet", "rename_sheet",
                                                              "number_format", "column_width", "delete_rows", "clear"]},
                            "sheet": {"type": ["string", "integer"]},
                            "cells": {"type": "object"}, "rows": {"type": "array", "items": {}}, "columns": {"type": "array", "items": {"type": "string"}},
                            "name": {"type": "string"}, "range": {"type": "string"}, "format": {"type": "string"},
                            "widths": {"type": "object"}, "row": {"type": "integer"}, "count": {"type": "integer"},
                        },
                        "required": ["op"],
                    },
                },
                "allow_chart_loss": {"type": "boolean", "description": "Only if the user agreed: the file has charts/pivot tables this tool cannot keep."},
            },
            "required": ["path", "ops"],
        },
    },
}


def _xl():
    openpyxl = O.need("openpyxl", "openpyxl")
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
    return openpyxl, Alignment, Border, Font, PatternFill, Side, column_index_from_string, get_column_letter, range_boundaries


# ------------------------------------------------------------------------------------------- values
_NUM = re.compile(r"^-?(?:0|[1-9]\d{0,14})(?:\.\d+)?$")
_NUM_COMMA = re.compile(r"^-?[1-9]\d{0,2}(?:,\d{3})+(?:\.\d+)?$")
_ISO_DATE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")


def clean_value(v: Any) -> Any:
    """Text that is really a number / date becomes one (so SUM and sorting work); everything else is kept as is."""
    if isinstance(v, str):
        s = v.strip()
        if s.startswith("="):
            return s
        if _NUM.match(s):
            return float(s) if "." in s else int(s)
        if _NUM_COMMA.match(s):
            n = s.replace(",", "")
            return float(n) if "." in n else int(n)
        m = _ISO_DATE.match(s)
        if m:
            try:
                return _dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            except ValueError:
                return v
    return v


def fmt_code(f: Any) -> str:
    s = str(f or "").strip()
    return FORMATS.get(s.lower(), s)


def _sheet_title(name: Any, taken: List[str]) -> str:
    t = re.sub(r"[\[\]:*?/\\]", " ", str(name or "").strip()).strip("' ")[:31] or "Sheet"
    base, n = t, 2
    while t.lower() in (x.lower() for x in taken):
        suffix = f" ({n})"
        t = base[:31 - len(suffix)] + suffix
        n += 1
    return t


def _norm_columns(cols: Any) -> Tuple[List[str], Dict[int, str], Dict[int, float]]:
    """['Name', {'name': 'Price', 'format': 'currency', 'width': 12}] -> headers, {col_idx: format}, {col_idx: width}"""
    heads: List[str] = []
    fmts: Dict[int, str] = {}
    widths: Dict[int, float] = {}
    for i, c in enumerate(O._as_list(cols, "columns") if cols not in (None, "", []) else [], 1):
        if isinstance(c, dict):
            heads.append(str(c.get("name", c.get("title", c.get("header", "")))))
            if c.get("format"):
                fmts[i] = fmt_code(c["format"])
            if c.get("width"):
                try:
                    widths[i] = float(c["width"])
                except (TypeError, ValueError):
                    pass
        else:
            heads.append("" if c is None else str(c))
    return heads, fmts, widths


def _norm_rows(rows: Any, heads: List[str]) -> Tuple[List[List[Any]], List[str]]:
    """Rows as lists; list-of-dicts rows also supply the headers when none were given."""
    lst = O._as_list(rows, "rows") if rows not in (None, "", []) else []
    if lst and all(isinstance(r, dict) for r in lst):
        if not heads:
            for r in lst:
                for k in r:
                    if k not in heads:
                        heads.append(str(k))
        lst = [[r.get(h, "") for h in heads] for r in lst]
    out = []
    for r in lst:
        out.append(list(r) if isinstance(r, (list, tuple)) else [r])
    return out, heads


def _range_cells(spec: str, ws, header_rows: int = 0):
    """'C' / 'C:C' / 'C2:C20' / 'B2' / 'A1:D9' -> (min_col, min_row, max_col, max_row) using the sheet size for open ends."""
    _, _, _, _, _, _, cidx, _, rb = _xl()
    s = str(spec or "").strip().upper().replace("$", "")
    if not s:
        raise O.OfficeError("a range is empty. Use e.g. C, C2:C20 or B2.")
    top = header_rows + 1
    if re.fullmatch(r"[A-Z]{1,3}", s):
        c = cidx(s)
        return c, top, c, max(ws.max_row, top)
    m = re.fullmatch(r"([A-Z]{1,3}):([A-Z]{1,3})", s)
    if m:
        return cidx(m.group(1)), top, cidx(m.group(2)), max(ws.max_row, top)
    try:
        c1, r1, c2, r2 = rb(s)
    except ValueError as e:
        raise O.OfficeError(f"cannot read the range '{spec}'. Use e.g. C, C2:C20 or B2.") from e
    return c1, r1, c2, r2


def write_sheet(wb, spec: Dict[str, Any], first: bool = False) -> Any:
    """Build one sheet from a spec dict (see _SHEET_PROPS). Returns the worksheet."""
    _, Alignment, Border, Font, PatternFill, Side, cidx, col_letter, _ = _xl()
    ws = wb.active if first else wb.create_sheet()
    ws.title = _sheet_title(spec.get("name"), [s.title for s in wb.worksheets if s is not ws])
    heads, col_fmts, col_widths = _norm_columns(spec.get("columns"))
    rows, heads = _norm_rows(spec.get("rows", spec.get("data")), heads)
    if not heads and not rows and not spec.get("formulas"):
        raise O.OfficeError(f"sheet '{ws.title}' is empty: give 'columns' and 'rows'.")
    body = Font(name=FONT, size=10)
    thin = Side(style="thin", color="D9D9D9")
    hdr_rows = 1 if heads else 0
    r0 = 1
    if heads:
        for ci, h in enumerate(heads, 1):
            c = ws.cell(row=1, column=ci, value=h)
            c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=HEAD_FILL)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        r0 = 2
    ncols = max([len(heads)] + [len(r) for r in rows] + [1])
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row, 1):
            c = ws.cell(row=r0 + ri, column=ci, value=clean_value(v))
            c.font = body
            c.border = Border(bottom=thin)
            if isinstance(c.value, _dt.date):
                c.number_format = FORMATS["date"]
    formulas = spec.get("formulas") or {}
    if isinstance(formulas, dict):
        for ref, f in formulas.items():
            try:
                c = ws[str(ref).strip().upper()]
            except Exception as e:                                 # noqa: BLE001
                raise O.OfficeError(f"sheet '{ws.title}': '{ref}' is not a cell like D12.") from e
            f = str(f)
            c.value = f if f.startswith("=") else "=" + f
            c.font = body
            c.border = Border(bottom=thin)
            ncols = max(ncols, c.column)
    for ci, f in col_fmts.items():                                  # formats from the column definitions
        for r in range(r0, max(ws.max_row, r0) + 1):
            ws.cell(row=r, column=ci).number_format = f
    nf = spec.get("number_formats") or {}
    if isinstance(nf, dict):
        for rng, f in nf.items():
            c1, r1, c2, r2 = _range_cells(rng, ws, hdr_rows)
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    ws.cell(row=r, column=c).number_format = fmt_code(f)
    for ci in range(1, ncols + 1):                                  # widths: given, else fitted to the content
        letter = col_letter(ci)
        w = col_widths.get(ci)
        cw = spec.get("column_widths")
        if w is None and isinstance(cw, dict):
            w = next((v for k, v in cw.items() if str(k).strip().upper() == letter), None)
        elif w is None and isinstance(cw, list) and ci - 1 < len(cw):
            w = cw[ci - 1]
        if w is None:
            longest = 0
            for r in range(1, ws.max_row + 1):
                v = ws.cell(row=r, column=ci).value
                n = 10 if isinstance(v, str) and v.startswith("=") else len(str(v)) if v is not None else 0
                longest = max(longest, n)
            w = min(50, max(8, longest * 1.15 + 2))
        try:
            ws.column_dimensions[letter].width = float(w)
        except (TypeError, ValueError):
            pass
    fr = spec.get("freeze", True if heads else False)
    if fr is True and heads:
        ws.freeze_panes = "A2"
    elif isinstance(fr, str) and fr.strip() and fr.strip().lower() not in ("true", "false"):
        ws.freeze_panes = fr.strip().upper()
    elif isinstance(fr, str) and fr.strip().lower() == "true" and heads:
        ws.freeze_panes = "A2"
    if heads and rows and spec.get("autofilter", True):
        ws.auto_filter.ref = f"A1:{col_letter(len(heads))}{r0 + len(rows) - 1}"
    return ws


# ------------------------------------------------------------------------------------------- xlsx_create
def _sheets_arg(args: dict) -> List[Dict[str, Any]]:
    sheets = args.get("sheets")
    if sheets in (None, "", []):
        if args.get("rows") or args.get("columns"):                  # forgiving: one sheet given at top level
            sheets = [{"name": args.get("name") or "Sheet1", "columns": args.get("columns"), "rows": args.get("rows")}]
        else:
            raise O.OfficeError("'sheets' is missing. Example: [{\"name\": \"Budget\", \"columns\": [\"Item\", \"Cost\"], \"rows\": [[\"Rent\", 500]]}].")
    out = []
    for i, s in enumerate(O._as_list(sheets, "sheets"), 1):
        if not isinstance(s, dict):
            raise O.OfficeError(f"sheet {i}: each sheet must be an object with 'name', 'columns' and 'rows'.")
        out.append(s)
    return out


def _check_formulas(target: Path) -> List[str]:
    errs = O.formula_errors(target)
    if not errs:
        return []
    return ["Formula errors: " + "; ".join(errs) + "."]


def _create(args: dict, ctx: ToolContext) -> str:
    try:
        openpyxl = O.need("openpyxl", "openpyxl")
        target = O.target_path(args.get("path"), (".xlsx",), must_exist=False, what="an Excel file")
        sheets = _sheets_arg(args)
        wb = openpyxl.Workbook()
        for i, s in enumerate(sheets):
            write_sheet(wb, s, first=(i == 0))
        wb.properties.creator = "Lucy"
        replaced = O.save_atomic(target, lambda tmp: wb.save(tmp))
        problems = _check_formulas(target)
        return O.finish(target, f"{len(sheets)} sheet(s)" + (", replaced the old file" if replaced else ""), problems)
    except Exception as e:                                           # noqa: BLE001
        return O.fail(e)


# ------------------------------------------------------------------------------------------- xlsx_edit
def _pick_sheet(wb, ref: Any, where: str):
    if ref in (None, ""):
        return wb.worksheets[0]
    if isinstance(ref, int) or (isinstance(ref, str) and ref.strip().isdigit()):
        i = int(ref)
        if 1 <= i <= len(wb.worksheets):
            return wb.worksheets[i - 1]
    for ws in wb.worksheets:
        if ws.title == str(ref):
            return ws
    for ws in wb.worksheets:
        if ws.title.lower() == str(ref).strip().lower():
            return ws
    raise O.OfficeError(f"{where}: no sheet '{ref}'. The sheets are: {', '.join(w.title for w in wb.worksheets)}.")


def _has_charts(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as z:
            return any(n.startswith(("xl/charts/", "xl/pivotTables/", "xl/slicers/")) for n in z.namelist())
    except Exception:                                                # noqa: BLE001
        return False


def _apply_ops(wb, ops: list) -> List[str]:
    _, Alignment, Border, Font, PatternFill, Side, cidx, col_letter, _ = _xl()
    body = Font(name=FONT, size=10)
    done: List[str] = []
    for n, op in enumerate(ops, 1):
        where = f"op {n}"
        if not isinstance(op, dict):
            raise O.OfficeError(f"{where}: each op must be an object like {{\"op\": \"set_cells\", \"cells\": {{\"B2\": 5}}}}.")
        kind = str(op.get("op") or op.get("action") or "").strip().lower().replace("-", "_")
        kind = {"set": "set_cells", "set_cell": "set_cells", "update": "set_cells", "add_rows": "append_rows", "append": "append_rows",
                "add_row": "append_rows", "new_sheet": "add_sheet", "remove_sheet": "delete_sheet", "format": "number_format",
                "set_format": "number_format", "set_width": "column_width", "column_widths": "column_width", "remove_rows": "delete_rows"}.get(kind, kind)
        if kind == "set_cells":
            ws = _pick_sheet(wb, op.get("sheet"), where)
            cells = op.get("cells")
            if isinstance(cells, list) and all(isinstance(x, dict) for x in cells):          # [{"cell": "B2", "value": 5}]
                cells = {str(x.get("cell", x.get("ref", ""))): x.get("value") for x in cells}
            if not isinstance(cells, dict) or not cells:
                raise O.OfficeError(f"{where}: set_cells needs 'cells', e.g. {{\"B2\": 5, \"C2\": \"=B2*2\"}}.")
            for ref, v in cells.items():
                try:
                    c = ws[str(ref).strip().upper()]
                except Exception as e:                                   # noqa: BLE001
                    raise O.OfficeError(f"{where}: '{ref}' is not a cell like B2.") from e
                c.value = clean_value(v)
                if c.font is None or c.font.name != FONT:
                    c.font = Font(name=FONT, size=c.font.size if c.font and c.font.size else 10, bold=c.font.bold if c.font else False)
                if isinstance(c.value, _dt.date):
                    c.number_format = FORMATS["date"]
            done.append(f"{len(cells)} cell(s) set on {ws.title}")
        elif kind == "append_rows":
            ws = _pick_sheet(wb, op.get("sheet"), where)
            rows, _h = _norm_rows(op.get("rows"), [])
            if not rows:
                raise O.OfficeError(f"{where}: append_rows needs 'rows' (a list of lists).")
            last = ws.max_row if any(c.value is not None for r in ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row) for c in r) else ws.max_row - 1
            for ri, row in enumerate(rows, 1):
                for ci, v in enumerate(row, 1):
                    c = ws.cell(row=last + ri, column=ci, value=clean_value(v))
                    c.font = body
                    above = ws.cell(row=last, column=ci)
                    if last > 1 and above.number_format and above.number_format != "General":
                        c.number_format = above.number_format
            done.append(f"{len(rows)} row(s) added to {ws.title}")
        elif kind == "add_sheet":
            spec = {k: op.get(k) for k in ("name", "columns", "rows", "formulas", "number_formats", "column_widths", "freeze") if op.get(k) is not None}
            write_sheet(wb, spec)
            done.append("sheet added")
        elif kind == "delete_sheet":
            ws = _pick_sheet(wb, op.get("sheet", op.get("name")), where)
            if len(wb.worksheets) == 1:
                raise O.OfficeError(f"{where}: a workbook needs at least one sheet.")
            wb.remove(ws)
            done.append(f"sheet {ws.title} deleted")
        elif kind == "rename_sheet":
            ws = _pick_sheet(wb, op.get("sheet"), where)
            new = _sheet_title(op.get("name", op.get("new_name")), [w.title for w in wb.worksheets if w is not ws])
            ws.title = new
            done.append(f"renamed to {new}")
        elif kind == "number_format":
            ws = _pick_sheet(wb, op.get("sheet"), where)
            if not op.get("format"):
                raise O.OfficeError(f"{where}: number_format needs 'format' (currency, percent, integer, number, date or an Excel code).")
            c1, r1, c2, r2 = _range_cells(op.get("range", op.get("cells")), ws, 1 if ws.max_row > 1 else 0)
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    ws.cell(row=r, column=c).number_format = fmt_code(op["format"])
            done.append(f"format set on {op.get('range')}")
        elif kind == "column_width":
            ws = _pick_sheet(wb, op.get("sheet"), where)
            widths = op.get("widths", op.get("width"))
            if not isinstance(widths, dict) or not widths:
                raise O.OfficeError(f"{where}: column_width needs 'widths', e.g. {{\"A\": 24, \"B\": 12}}.")
            for k, w in widths.items():
                try:
                    ws.column_dimensions[str(k).strip().upper()].width = float(w)
                except (TypeError, ValueError) as e:
                    raise O.OfficeError(f"{where}: width for {k} must be a number.") from e
            done.append("widths set")
        elif kind == "delete_rows":
            ws = _pick_sheet(wb, op.get("sheet"), where)
            try:
                row, count = int(op.get("row")), int(op.get("count", 1) or 1)
            except (TypeError, ValueError):
                raise O.OfficeError(f"{where}: delete_rows needs a numeric 'row' (from 1) and optional 'count'.")
            if row < 1 or count < 1 or row > ws.max_row:
                raise O.OfficeError(f"{where}: {ws.title} has {ws.max_row} row(s).")
            ws.delete_rows(row, count)
            done.append(f"{count} row(s) deleted (formulas below are NOT rewritten: check them)")
        elif kind == "clear":
            ws = _pick_sheet(wb, op.get("sheet"), where)
            c1, r1, c2, r2 = _range_cells(op.get("range"), ws, 0)
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    ws.cell(row=r, column=c).value = None
            done.append(f"{op.get('range')} cleared")
        else:
            raise O.OfficeError(f"{where}: unknown op '{kind}'. Use set_cells, append_rows, add_sheet, delete_sheet, rename_sheet, "
                                "number_format, column_width, delete_rows or clear.")
    return done


def _edit(args: dict, ctx: ToolContext) -> str:
    try:
        openpyxl = O.need("openpyxl", "openpyxl")
        target = O.target_path(args.get("path"), (".xlsx", ".xlsm"), must_exist=True, what="an Excel file")
        ops = O._as_list(args.get("ops"), "ops") if args.get("ops") not in (None, "", []) else []
        if not ops:
            raise O.OfficeError("'ops' is empty. Example: [{\"op\": \"set_cells\", \"sheet\": 1, \"cells\": {\"B2\": 500}}].")
        if _has_charts(target) and not args.get("allow_chart_loss"):
            raise O.OfficeError(f"{target.name} has charts or pivot tables, and this tool would drop them when saving. Nothing was changed. "
                                "Tell the user; if they agree to lose them, call again with allow_chart_loss=true.")
        try:
            wb = openpyxl.load_workbook(target, keep_vba=target.suffix.lower() == ".xlsm")
        except Exception as e:                                       # noqa: BLE001
            raise O.OfficeError(f"{target.name} could not be opened as an Excel file ({type(e).__name__}).") from e
        done = _apply_ops(wb, ops)
        O.save_atomic(target, lambda tmp: wb.save(tmp))
        return O.finish(target, "; ".join(done), _check_formulas(target))
    except Exception as e:                                           # noqa: BLE001
        return O.fail(e)


TOOLS = [
    {"schema": SCHEMA_CREATE, "handler": _create, "tier": "write", "label": "write"},
    {"schema": SCHEMA_EDIT, "handler": _edit, "tier": "write", "label": "edit"},
]
